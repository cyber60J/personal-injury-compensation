"""Conservative extraction: an amount is usable only with explicit dimensions.

``verified`` here means that extraction checks passed, not that a human approved
the source or the legal applicability of the statistic. Ambiguous candidates and
conflicts are retained for that separate review. No external requests are made.
"""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re
import unicodedata
import zipfile
from xml.etree import ElementTree

EXTRACTOR_VERSION = "statistical-evidence/1.0"
MAX_BYTES = 20 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_ROWS = 10000
MAX_COLUMNS = 200
MAX_CELLS = 200000
MAX_PAGES = 500

INCOME = "城镇居民人均可支配收入"
CONSUMPTION = "城镇居民人均消费支出"
NON_PRIVATE = "城镇非私营单位年平均工资"
PRIVATE = "城镇私营单位年平均工资"
SERVICE = "居民服务业年平均工资"
ON_DUTY = "城镇单位在岗职工年平均工资"

_REGIONS = (
    "北京", "天津", "上海", "重庆", "河北", "山西", "辽宁", "吉林", "黑龙江",
    "江苏", "浙江", "安徽", "福建", "江西", "山东", "河南", "湖北", "湖南",
    "广东", "海南", "四川", "贵州", "云南", "陕西", "甘肃", "青海", "内蒙古",
    "广西", "西藏", "宁夏", "新疆", "深圳", "厦门", "宁波", "青岛", "大连", "全国",
)
_YEAR = re.compile(r"(?<!\d)(20\d{2})(?:年(?:度)?)?(?!\d)")
_NUMBER = r"(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?"
_AMOUNT = re.compile(rf"^\s*({_NUMBER})\s*(万元|千元|百元|元)?\s*(?:[/／]人)?\s*$")
_INLINE_AMOUNT = re.compile(
    rf"^\s*(?:(20\d{{2}})年(?:度)?)?\s*(?:为人民币|约为|达到|为|达|是|约)?\s*({_NUMBER})\s*(万元|千元|百元|元)?(?![\d.%％])"
)
_UNIT = re.compile(r"万元|千元|百元|元")
_METRIC = re.compile(
    r"(?:城镇(?:常住)?居民)?人均(?:可支配收入|消费支出)|"
    r"城镇(?:常住)?居民(?:可支配收入|消费支出)|"
    r"(?:年)?平均工资"
)
_BAD_VALUE_COLUMN = re.compile(r"增长|增幅|增速|同比|比上年|增加额|增量|指数|占比|比重|%|％")
_FACTORS = {"元": 1, "百元": 100, "千元": 1000, "万元": 10000}
_OTHER_INDUSTRIES = (
    "农、林、牧、渔业", "农林牧渔业", "采矿业", "制造业", "电力、热力、燃气及水生产和供应业",
    "建筑业", "批发和零售业", "交通运输、仓储和邮政业", "住宿和餐饮业", "信息传输、软件和信息技术服务业",
    "金融业", "房地产业", "租赁和商务服务业", "科学研究和技术服务业", "水利、环境和公共设施管理业",
    "教育", "卫生和社会工作", "文化、体育和娱乐业", "公共管理、社会保障和社会组织",
)


def _clean(value: object) -> str:
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", str(value or "")))


def _years(value: str) -> set[int]:
    return {int(match.group(1)) for match in _YEAR.finditer(value)}


def _region_key(region: str) -> str:
    return re.sub(r"(?:壮族自治区|回族自治区|维吾尔自治区|自治区|省|市)$", "", region)


def _regions(value: str) -> set[str]:
    found = {region for region in _REGIONS if region in value}
    # Also catch a locally scoped city/county in a title/row instead of treating a
    # provincial site's identity as evidence for province-wide statistics.
    for match in re.finditer(r"(?:^|[\s，,。；;:：|]|20\d{2}年(?:度)?)([\u4e00-\u9fff]{2,12}(?:市|县|区|盟|州))(?=的|20\d{2}|城镇|居民|人均|统计|单位|$)", value):
        locality = _region_key(match.group(1))
        if locality not in {"全", "本", "该", "全地区", "本地区", "该地区"}:
            found.add(locality)
    return found


def _unit(value: str) -> str | None:
    found = set(_UNIT.findall(value))
    return next(iter(found)) if len(found) == 1 else None


def _metric_scope(local: str, enclosing: str, dimension_context: str = "") -> tuple[str, dict, list[str]] | None:
    """Only enclosing headings may complete a row's otherwise missing scope."""
    text = _clean(local)
    parent = _clean(enclosing)
    combined = text + "|" + parent
    reasons: list[str] = []
    metric = None
    if "可支配收入" in text:
        metric, field = "人均可支配收入", INCOME
    elif "消费支出" in text:
        metric, field = "人均消费支出", CONSUMPTION
    if metric:
        population_text = text if any(word in text for word in ("城镇居民", "城镇常住居民", "农村居民", "农村常住居民", "全体居民")) else parent
        if "农村" in population_text or "全体居民" in population_text:
            return None
        if not any(word in population_text for word in ("城镇居民", "城镇常住居民")):
            reasons.append("missing_urban_population")
        if "人均" not in combined:
            reasons.append("missing_per_capita_scope")
        return field, {"metric": metric, "population": "城镇居民" if not reasons else "未确认"}, reasons

    if "平均工资" not in combined or not any(word in text for word in ("平均工资", "居民服务", "就业人员", "在岗职工", "私营单位")):
        return None
    ownership_text = text if "私营" in text else parent
    non_private = "非私营" in ownership_text
    private = bool(re.search(r"(?<!非)私营", ownership_text))
    if non_private and private:
        reasons.append("conflicting_ownership")
    employment_text = text if "就业人员" in text or "在岗职工" in text else parent
    employment = "在岗职工" if "在岗职工" in employment_text else "就业人员" if "就业人员" in employment_text else "未确认"
    if "就业人员" in employment_text and "在岗职工" in employment_text:
        reasons.append("conflicting_employment_scope")
    industry = "居民服务、修理和其他服务业" if "居民服务" in text else None
    other_industries = [name for name in _OTHER_INDUSTRIES if _clean(name) in _clean(combined + dimension_context)]
    if industry:
        field = SERVICE
        if "修理" not in text or "其他服务业" not in text:
            reasons.append("incomplete_industry_scope")
        if not non_private or private:
            reasons.append("service_requires_non_private_units")
        if employment != "就业人员":
            reasons.append("service_requires_employed_persons")
    elif employment == "在岗职工":
        field = ON_DUTY
        if non_private or private:
            reasons.append("on_duty_ownership_subgroup_not_total")
    elif non_private:
        field = NON_PRIVATE
    elif private:
        field = PRIVATE
    else:
        return None
    if employment == "未确认":
        reasons.append("missing_employment_scope")
    if other_industries:
        reasons.append("non_total_industry_scope")
        industry = "、".join(other_industries)
    if any(term in combined for term in ("规模以上", "国有单位", "集体单位", "外商投资", "港澳台投资", "事业单位", "机关单位")):
        reasons.append("restricted_unit_subgroup")
    if "城镇" not in combined:
        reasons.append("missing_urban_units_scope")
    if "年平均工资" not in combined:
        reasons.append("missing_annual_wage_scope")
    return field, {
        "metric": "年平均工资", "population": "城镇单位", "employment": employment,
        "ownership": "非私营" if non_private and not private else "私营" if private and not non_private else "全部或未确认",
        "industry": industry or "全部行业",
    }, reasons


@dataclass
class _Cell:
    text: str
    locator: dict


class _Extraction:
    def __init__(self, region: str, year: int, title: str, source_url: str):
        self.region, self.year = region, year
        self.title, self.source_url = title, source_url
        self.candidates: dict[str, list[dict]] = {}
        self.warnings: list[str] = []

    def add(self, metric: tuple[str, dict, list[str]], raw_value: str, raw_unit: str | None,
            *, year_contexts: list[str], region_contexts: list[str], snippet: str,
            locator: dict, method: str, extra_reasons: list[str] | None = None):
        field, scope, initial_reasons = metric
        reasons = list(initial_reasons) + list(extra_reasons or [])
        inferred_year = None
        for context in year_contexts:
            found = _years(context)
            if found:
                if len(found) == 1:
                    inferred_year = next(iter(found))
                else:
                    reasons.append("ambiguous_statistical_year")
                break
        if inferred_year is None and "ambiguous_statistical_year" not in reasons:
            reasons.append("missing_statistical_year")
        if inferred_year is not None and inferred_year != self.year:
            reasons.append("statistical_year_mismatch")
        inferred_region = None
        target_region = _region_key(self.region)
        for context in region_contexts:
            found_regions = _regions(context)
            if found_regions:
                if found_regions == {target_region}:
                    inferred_region = self.region
                elif len(found_regions) == 1:
                    inferred_region = next(iter(found_regions))
                    reasons.append("region_mismatch")
                else:
                    reasons.append("ambiguous_region")
                break
        if inferred_region is None and "ambiguous_region" not in reasons:
            reasons.append("missing_region")
        if raw_unit is None:
            reasons.append("missing_or_ambiguous_unit")
        try:
            raw_number = Decimal(_clean(raw_value).replace(",", ""))
            number = raw_number * _FACTORS.get(raw_unit or "", 1)
            value = float(number)
        except (ValueError, InvalidOperation, OverflowError):
            return
        if not number.is_finite() or number <= 0 or number > 10000000:
            reasons.append("amount_out_of_range")
        candidate = {
            "value": value, "raw_value": raw_value, "raw_unit": raw_unit,
            "unit": "元" if raw_unit else None,
            "conversion": {"factor": _FACTORS[raw_unit], "operation": f"{raw_value} × {_FACTORS[raw_unit]}"} if raw_unit else None,
            "region": inferred_region, "statistical_year": inferred_year, "scope": scope,
            "snippet": snippet[:2500], "locator": locator,
            "extract_method": method, "extractor_version": EXTRACTOR_VERSION,
            "source_url": self.source_url, "document_title": self.title,
            "validation_status": "pending" if reasons else "verified",
            "review_reasons": sorted(set(reasons)),
        }
        self.candidates.setdefault(field, []).append(candidate)

    def paragraphs(self, paragraphs: list[tuple[str, dict]], method: str):
        for text, locator in paragraphs:
            # A sentence is the maximum context for scope. This prevents the
            # previous sentence's wage population from leaking into a new one.
            for sentence in re.split(r"[。；;\n]", unicodedata.normalize("NFKC", text)):
                for match in _METRIC.finditer(sentence):
                    amount = _INLINE_AMOUNT.match(sentence[match.end():])
                    if amount is None:
                        continue
                    prefix = sentence[:match.end()]
                    # Use this clause's label, plus its own sentence heading;
                    # unrelated metrics before a comma must not supply scope.
                    clause_prefix = re.split(r"[，,]", prefix)[-1]
                    metric = _metric_scope(clause_prefix, self.title, prefix)
                    if metric is None:
                        continue
                    year_prefix = prefix
                    found_years = list(_YEAR.finditer(year_prefix))
                    local_year = amount.group(1) or (found_years[-1].group(1) if found_years else "")
                    extra_reasons: list[str] = []
                    relative_years = list(re.finditer(r"(?<!比)(?:上一年|上年度|上年|去年)", prefix))
                    if relative_years and (not found_years or relative_years[-1].start() > found_years[-1].end()):
                        base_years = {int(local_year)} if local_year else _years(self.title)
                        if len(base_years) == 1:
                            local_year = str(next(iter(base_years)) - 1)
                        else:
                            extra_reasons.append("relative_year_without_base")
                    # Unit-less prose is retained for review; percentages are
                    # not monetary candidates, even when followed by a total.
                    after = sentence[match.end() + amount.end():].lstrip()
                    if after.startswith(("%", "％")):
                        continue
                    self.add(metric, amount.group(2), amount.group(3),
                             year_contexts=[local_year, self.title],
                             region_contexts=[prefix, self.title], snippet=sentence,
                             locator=locator, method=method, extra_reasons=extra_reasons)

    def table(self, grid: list[list[_Cell | None]], context: str, method: str, footnotes: list[str] | None = None):
        visited: set[tuple] = set()
        column_headers: dict[int, dict[str, None]] = {}
        for row_number, row in enumerate(grid):
            for column_number, cell in enumerate(row):
                if cell is None:
                    continue
                identity = tuple(sorted((str(k), str(v)) for k, v in cell.locator.items()))
                if identity in visited:
                    continue
                amount = _AMOUNT.fullmatch(_clean(cell.text))
                if amount is None or re.fullmatch(r"20\d{2}(?:年(?:度)?)?", _clean(cell.text)):
                    continue
                visited.add(identity)
                left = [item.text for item in row[:column_number] if item is not None and not _is_amount(item.text)]
                above = list(column_headers.get(column_number, {}))
                # A year may itself be a numeric cell. Keep years as dimensions,
                # while excluding previous rows' unrelated monetary values.
                left += [item.text for item in row[:column_number] if item and re.fullmatch(r"20\d{2}(?:年(?:度)?)?", _clean(item.text))]
                left_text, above_text = " | ".join(dict.fromkeys(left)), " | ".join(dict.fromkeys(above))
                if _BAD_VALUE_COLUMN.search(above_text) or _BAD_VALUE_COLUMN.search(left_text):
                    continue
                # Table headings (including spanning rows) can provide units
                # and population, but unrelated indicator rows cannot.
                local = left_text + " | " + above_text
                metric = _metric_scope(local, context + " | " + self.title)
                if metric is None:
                    continue
                raw_unit = amount.group(2) or _unit(above_text) or _unit(left_text) or _unit(context)
                # Explicit period cells outrank a spanning document heading.
                # This also supports tables where statistical years are rows.
                column_period = " | ".join(value for value in above if re.fullmatch(r"20\d{2}(?:年(?:度)?)?", _clean(value)))
                row_period = " | ".join(value for value in left if re.fullmatch(r"20\d{2}(?:年(?:度)?)?", _clean(value)))
                self.add(metric, amount.group(1), raw_unit,
                         year_contexts=[column_period, row_period, above_text, left_text, context, self.title],
                         region_contexts=[left_text, above_text, context, self.title],
                         snippet=f"{context}\n行: {left_text}\n列: {above_text}\n单元格: {cell.text}",
                         locator=dict(cell.locator, row_label=left_text, column_label=above_text,
                                      headers=list(dict.fromkeys(left + above)), footnotes=footnotes or []),
                         method=method)
            for column_number, cell in enumerate(row):
                if cell and cell.text and (not _is_amount(cell.text) or re.fullmatch(r"20\d{2}(?:年(?:度)?)?", _clean(cell.text))):
                    column_headers.setdefault(column_number, {})[cell.text] = None

    def result(self) -> dict:
        data, evidence = {}, {}
        for field, candidates in self.candidates.items():
            verified = [candidate for candidate in candidates if candidate["validation_status"] == "verified"]
            unique_values = {candidate["value"] for candidate in verified}
            if len(unique_values) == 1:
                data[field] = verified[0]["value"]
                evidence[field] = dict(verified[0], corroborating_locators=[candidate["locator"] for candidate in verified[1:]])
            elif len(unique_values) > 1:
                for candidate in verified:
                    candidate["validation_status"] = "pending"
                    candidate["review_reasons"].append("conflicting_values")
                evidence[field] = dict(verified[0], conflicts=verified)
                self.warnings.append(f"conflicting_values: {field}")
            else:
                evidence[field] = dict(candidates[0])
        return {"data": data, "evidence": evidence, "candidates": self.candidates,
                "warnings": list(dict.fromkeys(self.warnings))}


def _is_amount(text: str) -> bool:
    return bool(_AMOUNT.fullmatch(_clean(text)))


def _html(content: bytes, extraction: _Extraction):
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(content, "html.parser")
    for element in soup(["script", "style", "nav", "footer"]):
        element.decompose()
    heading = soup.find("h1") or soup.find("title")
    if heading and _years(heading.get_text(" ", strip=True)):
        extraction.title = heading.get_text(" ", strip=True)
    for table_index, table in enumerate(soup.find_all("table"), 1):
        if table.find_parent("table"):
            extraction.warnings.append("nested_table_requires_review")
            continue
        context_parts = []
        caption = table.find("caption")
        if caption:
            context_parts.append(caption.get_text(" ", strip=True))
        for sibling in list(table.previous_siblings)[:6]:
            if getattr(sibling, "name", None) == "table":
                break
            value = sibling.get_text(" ", strip=True) if hasattr(sibling, "get_text") else str(sibling).strip()
            if value and len(value) < 300:
                context_parts.append(value)
        grid: list[list[_Cell | None]] = []
        count = 0
        for row_number, row in enumerate(table.find_all("tr", recursive=True)):
            if row.find_parent("table") is not table:
                continue
            if row_number >= MAX_ROWS:
                raise ValueError("table_row_limit")
            while len(grid) <= row_number:
                grid.append([])
            column_number = 0
            for item in row.find_all(["td", "th"], recursive=False):
                while column_number < len(grid[row_number]) and grid[row_number][column_number] is not None:
                    column_number += 1
                try:
                    rowspan, colspan = int(item.get("rowspan", 1)), int(item.get("colspan", 1))
                except ValueError:
                    raise ValueError("invalid_table_span") from None
                if not (1 <= rowspan <= MAX_ROWS and 1 <= colspan <= MAX_COLUMNS) or column_number + colspan > MAX_COLUMNS or row_number + rowspan > MAX_ROWS:
                    raise ValueError("table_span_limit")
                cell = _Cell(item.get_text(" ", strip=True), {"type": "html_table", "table": table_index,
                             "row": row_number + 1, "column": column_number + 1,
                             "rowspan": rowspan, "colspan": colspan})
                count += rowspan * colspan
                if count > MAX_CELLS:
                    raise ValueError("table_cell_limit")
                for r in range(row_number, row_number + rowspan):
                    while len(grid) <= r:
                        grid.append([])
                    while len(grid[r]) < column_number + colspan:
                        grid[r].append(None)
                    for c in range(column_number, column_number + colspan):
                        grid[r][c] = cell
                column_number += colspan
        footnotes = [element.get_text(" ", strip=True) for element in table.find_all("tfoot")]
        for sibling in list(table.next_siblings)[:4]:
            value = sibling.get_text(" ", strip=True) if hasattr(sibling, "get_text") else str(sibling).strip()
            if value and re.match(r"(?:注[：:0-9一二三（(]|说明[：:])", value):
                footnotes.append(value[:2000])
            elif value:
                break
        extraction.table(grid, " | ".join(context_parts), "html_table", footnotes)
    for table in soup.find_all("table"):
        table.decompose()
    paragraphs = []
    for index, element in enumerate(soup.find_all(["p", "li"]), 1):
        paragraphs.append((element.get_text(" ", strip=True), {"type": "html_paragraph", "paragraph": index}))
        element.decompose()
    paragraphs += [(line, {"type": "html_text", "line": index})
                   for index, line in enumerate(soup.get_text("\n", strip=True).splitlines(), 1)]
    extraction.paragraphs(paragraphs, "html_text")


def _xlsx(content: bytes, extraction: _Extraction):
    from openpyxl import load_workbook
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, range_boundaries

    # read_only avoids loading worksheet objects; merge definitions are retained
    # separately from the XML because ReadOnlyWorksheet omits merged_cells.
    with zipfile.ZipFile(BytesIO(content)) as archive:
        if sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
            raise ValueError("xlsx_uncompressed_size_limit")
        namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        merges: dict[str, list[str]] = {}
        actual_cells = 0
        for name in archive.namelist():
            if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name):
                root = ElementTree.fromstring(archive.read(name))
                for row in root.findall("s:sheetData/s:row", namespace):
                    if int(row.get("r", "0")) > MAX_ROWS:
                        raise ValueError("xlsx_row_limit")
                    for cell in row.findall("s:c", namespace):
                        actual_cells += 1
                        if actual_cells > MAX_CELLS:
                            raise ValueError("xlsx_cell_limit")
                        column, row_index = coordinate_from_string(cell.attrib["r"])
                        if row_index > MAX_ROWS or column_index_from_string(column) > MAX_COLUMNS:
                            raise ValueError("xlsx_dimension_limit")
                merges[name] = [item.attrib["ref"] for item in root.findall("s:mergeCells/s:mergeCell", namespace)]
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        try:
            total = 0
            for sheet in workbook.worksheets:
                if (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_COLUMNS:
                    raise ValueError("xlsx_dimension_limit")
                total += (sheet.max_row or 0) * (sheet.max_column or 0)
                if total > MAX_CELLS:
                    raise ValueError("xlsx_cell_limit")
                grid = [[_Cell(str(cell.value) if cell.value is not None else "",
                               {"type": "xlsx_cell", "sheet": sheet.title,
                                "cell": getattr(cell, "coordinate", None)}) for cell in row]
                        for row in sheet.iter_rows()]
                path = str(getattr(sheet, "_worksheet_path", "")).lstrip("/")
                for merged in merges.get(path, []):
                    c1, r1, c2, r2 = range_boundaries(merged)
                    if r2 > MAX_ROWS or c2 > MAX_COLUMNS or (r2-r1+1)*(c2-c1+1) > MAX_CELLS:
                        raise ValueError("xlsx_merge_limit")
                    if r1 > len(grid) or c1 > len(grid[r1-1]):
                        continue
                    original = grid[r1-1][c1-1]
                    for r in range(r1-1, min(r2, len(grid))):
                        for c in range(c1-1, min(c2, len(grid[r]))):
                            grid[r][c] = original
                leading = []
                for row in grid[:8]:
                    texts = list(dict.fromkeys(cell.text for cell in row if cell and cell.text))
                    if len(texts) == 1 and not _is_amount(texts[0]):
                        leading.append(texts[0])
                    elif len(texts) > 1:
                        break
                context = " | ".join([sheet.title] + leading)
                extraction.table(grid, context, "xlsx_table")
        finally:
            workbook.close()


def _pdf(content: bytes, extraction: _Extraction):
    import pdfplumber

    with pdfplumber.open(BytesIO(content)) as document:
        if len(document.pages) > MAX_PAGES:
            raise ValueError("pdf_page_limit")
        for page_number, page in enumerate(document.pages, 1):
            text = page.extract_text() or ""
            if not text.strip():
                extraction.warnings.append(f"needs_ocr: page {page_number}")
                continue
            if page_number == 1:
                # Read an explicit document heading, never infer a statistical
                # year from a publication date or from the body's first amount.
                lines = [line.strip() for line in text.splitlines() if line.strip()]
                for size in range(1, min(3, len(lines)) + 1):
                    heading = " ".join(lines[:size])
                    if len(heading) > 200 or re.search(rf"{_NUMBER}\s*(?:万元|元|%)", heading):
                        break
                    if _years(heading) and _regions(heading) and any(term in heading for term in ("统计公报", "平均工资", "收入和消费", "收入消费")):
                        extraction.title = heading
                        break
            tables = page.find_tables()
            for table_number, table in enumerate(tables, 1):
                rows = table.extract()
                if len(rows) > MAX_ROWS or any(len(row) > MAX_COLUMNS for row in rows) or sum(map(len, rows)) > MAX_CELLS:
                    raise ValueError("pdf_table_limit")
                grid = []
                for r, row in enumerate(rows):
                    cells = []
                    for c, value in enumerate(row):
                        bbox = table.rows[r].cells[c] if r < len(table.rows) and c < len(table.rows[r].cells) else table.bbox
                        cells.append(_Cell(value or "", {"type": "pdf_table", "page": page_number,
                                     "table": table_number, "row": r+1, "column": c+1,
                                     "bbox": list(bbox) if bbox else list(table.bbox)}))
                    grid.append(cells)
                # The visible heading immediately above the table supplies its
                # units; unrelated text beneath it cannot do so.
                top = max(0, float(table.bbox[1]) - 100)
                context = page.crop((0, top, page.width, table.bbox[1])).extract_text() or "" if table.bbox[1] > top else ""
                extraction.table(grid, context, "pdf_table")
            table_boxes = [table.bbox for table in tables]
            words = page.extract_words()
            lines: dict[float, list[dict]] = {}
            for word in words:
                if any(box[0] <= word["x0"] and word["x1"] <= box[2] and box[1] <= word["top"] and word["bottom"] <= box[3] for box in table_boxes):
                    continue
                lines.setdefault(round(word["top"], 0), []).append(word)
            paragraphs = []
            for _, line in sorted(lines.items()):
                line.sort(key=lambda item: item["x0"])
                paragraphs.append(("".join(word["text"] for word in line), {"type": "pdf_text", "page": page_number,
                                   "bbox": [min(w["x0"] for w in line), min(w["top"] for w in line),
                                            max(w["x1"] for w in line), max(w["bottom"] for w in line)]}))
            extraction.paragraphs(paragraphs, "pdf_text")


def extract_document(content: bytes, content_type: str, *, region: str, year: int,
                     title: str | None = None, source_url: str = "") -> dict:
    """Extract target statistics with provenance, retaining uncertain candidates.

    Only candidates with explicit compatible region, statistical year, monetary
    unit and population enter ``data``. Format failures return actionable warnings
    and never cause a fallback to proximity-based numeric extraction.
    """
    extraction = _Extraction(region, year, title or "", source_url)
    if len(content) > MAX_BYTES:
        extraction.warnings.append("document_size_limit")
        return extraction.result()
    try:
        if content.startswith(b"%PDF") or "pdf" in content_type.lower():
            _pdf(content, extraction)
        elif content.startswith(b"PK") or "spreadsheetml" in content_type.lower() or content_type.lower() == "xlsx":
            _xlsx(content, extraction)
        elif "html" in content_type.lower() or content.lstrip().startswith(b"<"):
            _html(content, extraction)
        elif content_type.lower().startswith("text/plain"):
            text = content.decode("utf-8-sig")
            extraction.paragraphs([(line, {"type": "text", "line": index})
                                   for index, line in enumerate(text.splitlines(), 1)], "plain_text")
        else:
            extraction.warnings.append("unsupported_document_format")
    except Exception as error:
        # Discard partial work after a format/size failure: it may omit a later
        # conflicting amount or silently drop dimensions from damaged content.
        extraction.candidates.clear()
        extraction.warnings.append(f"document_extraction_failed: {type(error).__name__}: {error}")
    return extraction.result()
