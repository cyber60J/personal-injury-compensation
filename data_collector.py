#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
全国人身损害赔偿核心统计数据收集脚本

功能：自动收集全国31个省、自治区、直辖市 + 5个计划单列市（深圳、厦门、宁波、青岛、大连）
      过去10个完整统计年度的6项核心统计数据

数据项：
1. 城镇居民人均可支配收入
2. 城镇居民人均消费支出
3. 城镇非私营单位就业人员年平均工资
4. 城镇私营单位就业人员年平均工资
5. 居民服务业年平均工资
6. 城镇单位在岗职工年平均工资

数据源：各地区统计局官网发布的统计公报
"""

from __future__ import annotations

import argparse
import ipaddress
import os
import re
import time
import json
import logging
import hashlib
import sys
from pathlib import Path
from decimal import Decimal
from collections import OrderedDict
from typing import Dict, List, Tuple, Optional, Any
from dataclasses import dataclass, asdict, field
from datetime import datetime
import urllib.parse

# Keep the standalone collector usable before an editable package installation.
_src_dir = Path(__file__).resolve().parent / "src"
if _src_dir.is_dir() and str(_src_dir) not in sys.path:
    sys.path.insert(0, str(_src_dir))

from personal_injury.collection.source_archive import ArchiveStore

MISSING_DEPENDENCIES: List[str] = []


def _record_missing_dependency(error: ModuleNotFoundError, fallback_name: str):
    package_names = {
        "bs4": "beautifulsoup4",
    }
    missing_name = package_names.get(error.name or "", fallback_name)
    if missing_name not in MISSING_DEPENDENCIES:
        MISSING_DEPENDENCIES.append(missing_name)


try:
    import requests
    from requests.adapters import HTTPAdapter
    from urllib3.util.retry import Retry
except ModuleNotFoundError as exc:
    requests = None
    HTTPAdapter = None
    Retry = None
    _record_missing_dependency(exc, "requests")

try:
    from bs4 import BeautifulSoup
except ModuleNotFoundError as exc:
    BeautifulSoup = None
    _record_missing_dependency(exc, "beautifulsoup4")

try:
    import pandas as pd
except ModuleNotFoundError as exc:
    pd = None
    _record_missing_dependency(exc, "pandas")

try:
    from openpyxl.styles import PatternFill, Font, Alignment
except ModuleNotFoundError as exc:
    PatternFill = None
    Font = None
    Alignment = None
    _record_missing_dependency(exc, "openpyxl")

try:
    import chardet
except ModuleNotFoundError:
    chardet = None

# ==================== 配置区域 ====================
# 可配置参数，便于后续修改

# 统计年份范围：默认最近10个完整统计年度。
# 例如在2026年运行时，默认收集2016-2025年。
END_YEAR = datetime.now().year - 1
START_YEAR = END_YEAR - 9

# 地区列表：大陆31个省级行政区 + 5个计划单列市。
# 不把缺少可核验官方统计局URL的地区放入默认采集范围，避免制造必然失败的数据行。
REGIONS = [
    # 直辖市
    "北京市", "天津市", "上海市", "重庆市",
    # 省
    "河北省", "山西省", "辽宁省", "吉林省", "黑龙江省",
    "江苏省", "浙江省", "安徽省", "福建省", "江西省",
    "山东省", "河南省", "湖北省", "湖南省", "广东省",
    "海南省", "四川省", "贵州省", "云南省", "陕西省",
    "甘肃省", "青海省",
    # 自治区
    "内蒙古自治区", "广西壮族自治区", "西藏自治区",
    "宁夏回族自治区", "新疆维吾尔自治区",
    # 计划单列市
    "深圳市", "厦门市", "宁波市", "青岛市", "大连市"
]

# 各地区统计局官网URL（需要根据实际情况补充或更新）
# 格式：{"地区名称": "统计局官网基础URL"}
REGION_STATS_URLS = {
    "北京市": "http://tjj.beijing.gov.cn",
    "天津市": "http://stats.tj.gov.cn",
    "上海市": "https://tjj.sh.gov.cn",
    "重庆市": "http://tjj.cq.gov.cn",
    "河北省": "http://tjj.hebei.gov.cn",
    "山西省": "http://tjj.shanxi.gov.cn",
    "辽宁省": "http://tjj.ln.gov.cn",
    "吉林省": "http://tjj.jl.gov.cn",
    "黑龙江省": "http://tjj.hlj.gov.cn",
    "江苏省": "http://tj.jiangsu.gov.cn",
    "浙江省": "http://tjj.zj.gov.cn",
    "安徽省": "http://tjj.ah.gov.cn",
    "福建省": "http://tjj.fujian.gov.cn",
    "江西省": "http://tjj.jiangxi.gov.cn",
    "山东省": "http://tjj.shandong.gov.cn",
    "河南省": "http://tjj.henan.gov.cn",
    "湖北省": "http://tjj.hubei.gov.cn",
    "湖南省": "http://tjj.hunan.gov.cn",
    "广东省": "http://stats.gd.gov.cn",
    "海南省": "http://stats.hainan.gov.cn",
    "四川省": "http://tjj.sc.gov.cn",
    "贵州省": "http://tjj.guizhou.gov.cn",
    "云南省": "http://stats.yn.gov.cn",
    "陕西省": "http://tjj.shaanxi.gov.cn",
    "甘肃省": "http://tjj.gansu.gov.cn",
    "青海省": "http://tjj.qinghai.gov.cn",
    "内蒙古自治区": "http://tjj.nmg.gov.cn",
    "广西壮族自治区": "http://tjj.gxzf.gov.cn",
    "西藏自治区": "http://tjj.xizang.gov.cn",
    "宁夏回族自治区": "http://tjj.nx.gov.cn",
    "新疆维吾尔自治区": "http://tjj.xinjiang.gov.cn",
    # 计划单列市
    "深圳市": "http://tjj.sz.gov.cn",
    "厦门市": "http://tjj.xm.gov.cn",
    "宁波市": "http://tjj.ningbo.gov.cn",
    "青岛市": "http://tjj.qingdao.gov.cn",
    "大连市": "http://tjj.dl.gov.cn",
}

# 需要搜索的公报类型及关键词（按优先级排序）
BULLETIN_TYPES = [
    {
        "name": "国民经济和社会发展统计公报",
        "keywords": ["国民经济和社会发展统计公报", "统计公报", "经济社会统计公报"]
    },
    {
        "name": "城镇单位就业人员平均工资公报",
        "keywords": [
            "城镇单位就业人员年平均工资",
            "城镇单位就业人员平均工资",
            "城镇非私营单位就业人员年平均工资",
            "城镇私营单位就业人员年平均工资",
            "就业人员年平均工资情况",
            "就业人员平均工资公报",
            "平均工资情况",
            "平均工资公报"
        ]
    },
    {
        "name": "居民收入和消费支出情况公报",
        "keywords": ["居民收入和消费支出", "居民收入消费支出公报", "收入消费支出情况"]
    }
]

# 数据项定义及对应的搜索关键词
DATA_ITEMS = {
    "城镇居民人均可支配收入": [
        "城镇居民人均可支配收入",
        "城镇常住居民人均可支配收入",
        "城镇居民可支配收入"
    ],
    "城镇居民人均消费支出": [
        "城镇居民人均消费支出",
        "城镇常住居民人均消费支出",
        "城镇居民消费支出"
    ],
    "城镇非私营单位年平均工资": [
        "城镇非私营单位就业人员年平均工资",
    ],
    "城镇私营单位年平均工资": [
        "城镇私营单位就业人员年平均工资",
    ],
    "居民服务业年平均工资": [
        "居民服务、修理和其他服务业就业人员年平均工资",
        "居民服务业年平均工资",
        "居民服务业平均工资"
    ],
    "城镇单位在岗职工年平均工资": [
        "城镇单位在岗职工年平均工资",
        "城镇在岗职工年平均工资",
        "在岗职工年平均工资"
    ]
}

DATA_FIELD_ATTRS = {
    "城镇居民人均可支配收入": "urban_disposable_income",
    "城镇居民人均消费支出": "urban_consumption_expenditure",
    "城镇非私营单位年平均工资": "urban_non_private_wage",
    "城镇私营单位年平均工资": "urban_private_wage",
    "居民服务业年平均工资": "service_industry_wage",
    "城镇单位在岗职工年平均工资": "urban_on_duty_wage",
}

DATA_FIELD_COLUMNS = {
    "城镇居民人均可支配收入": "城镇居民人均可支配收入(元)",
    "城镇居民人均消费支出": "城镇居民人均消费支出(元)",
    "城镇非私营单位年平均工资": "城镇非私营单位年平均工资(元)",
    "城镇私营单位年平均工资": "城镇私营单位年平均工资(元)",
    "居民服务业年平均工资": "居民服务等行业非私营就业人员年平均工资(元)",
    "城镇单位在岗职工年平均工资": "城镇单位在岗职工年平均工资(元)",
}

REQUIRED_DATA_ITEMS = [
    "城镇居民人均可支配收入",
    "城镇居民人均消费支出",
    "城镇非私营单位年平均工资",
    "城镇私营单位年平均工资",
    "居民服务业年平均工资",
]

OPTIONAL_LEGACY_DATA_ITEMS = [
    "城镇单位在岗职工年平均工资",
]

VALUE_RANGE_RULES = {
    "城镇居民人均可支配收入": (5000, 200000),
    "城镇居民人均消费支出": (3000, 150000),
    "城镇非私营单位年平均工资": (10000, 500000),
    "城镇私营单位年平均工资": (5000, 300000),
    "居民服务业年平均工资": (5000, 300000),
    "城镇单位在岗职工年平均工资": (10000, 500000),
}

# 网络请求配置
REQUEST_TIMEOUT = 30  # 秒
MAX_RETRIES = 3
RETRY_BACKOFF_FACTOR = 1.5
MAX_REDIRECTS = 5
MAX_RESPONSE_BYTES = 20 * 1024 * 1024
MAX_PDF_PAGES = 500
REDIRECT_STATUS_CODES = {301, 302, 303, 307, 308}


def is_allowed_official_url(candidate_url: str, official_url: str) -> bool:
    """仅允许同一官方域名体系内的HTTP(S)地址。"""
    try:
        candidate = urllib.parse.urlsplit(candidate_url)
        official = urllib.parse.urlsplit(official_url)
        candidate_host = (candidate.hostname or "").lower().rstrip(".")
        official_host = (official.hostname or "").lower().rstrip(".")
        candidate_port = candidate.port
        official_port = official.port
    except (TypeError, ValueError):
        return False

    if candidate.scheme not in {"http", "https"} or official.scheme not in {"http", "https"}:
        return False
    if not candidate_host or not official_host:
        return False
    if candidate.username or candidate.password or official.username or official.password:
        return False
    if candidate_port not in {None, 80, 443} or official_port not in {None, 80, 443}:
        return False
    if official.scheme == "https" and candidate.scheme != "https":
        return False
    if candidate_host != official_host and not candidate_host.endswith("." + official_host):
        return False
    if candidate_host == "localhost" or candidate_host.endswith(".localhost"):
        return False

    try:
        ipaddress.ip_address(candidate_host)
    except ValueError:
        return True
    return False


def fetch_official_content(
    session: requests.Session,
    url: str,
    official_url: str,
    *,
    timeout: int = REQUEST_TIMEOUT,
    max_bytes: int = MAX_RESPONSE_BYTES,
    method: str = "GET",
    json_body: Optional[Dict[str, Any]] = None,
) -> Tuple[Any, bytes]:
    """按跳校验重定向并限制下载体积，避免越域抓取和超大响应。"""
    method = method.upper()
    if method not in {"GET", "POST"}:
        raise ValueError("官方来源仅支持GET或POST读取")
    current_url = url
    for redirect_count in range(MAX_REDIRECTS + 1):
        if not is_allowed_official_url(current_url, official_url):
            raise ValueError(f"拒绝访问非官方域名URL：{current_url}")

        request_options = dict(timeout=timeout, allow_redirects=False, stream=True)
        if method == "POST":
            response = session.post(current_url, json=json_body, **request_options)
        else:
            response = session.get(current_url, **request_options)
        response_url = getattr(response, "url", None) or current_url
        if not is_allowed_official_url(response_url, official_url):
            response.close()
            raise ValueError(f"响应URL越出官方域名范围：{response_url}")

        if response.status_code in REDIRECT_STATUS_CODES:
            location = response.headers.get("Location")
            response.close()
            if method == "POST":
                raise ValueError("拒绝跟随官方数据POST请求的重定向")
            if not location:
                raise ValueError(f"重定向响应缺少Location：{response_url}")
            if redirect_count >= MAX_REDIRECTS:
                raise ValueError(f"重定向次数超过{MAX_REDIRECTS}次：{url}")
            current_url = urllib.parse.urljoin(response_url, location)
            continue

        content_length = response.headers.get("Content-Length")
        if content_length:
            try:
                if int(content_length) > max_bytes:
                    raise ValueError(f"响应体超过{max_bytes}字节上限：{response_url}")
            except ValueError as exc:
                response.close()
                if "超过" in str(exc):
                    raise

        chunks = []
        downloaded = 0
        try:
            for chunk in response.iter_content(chunk_size=64 * 1024):
                if not chunk:
                    continue
                downloaded += len(chunk)
                if downloaded > max_bytes:
                    raise ValueError(f"响应体超过{max_bytes}字节上限：{response_url}")
                chunks.append(chunk)
        finally:
            response.close()
        return response, b"".join(chunks)

    raise ValueError(f"无法完成官方来源请求：{url}")


def spreadsheet_safe_value(value: Any) -> Any:
    """防止外部网页文本在Excel或CSV中被解释为公式。"""
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def spreadsheet_safe_rows(rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    return [
        {column: spreadsheet_safe_value(value) for column, value in row.items()}
        for row in rows
    ]

# 输出文件路径
OUTPUT_EXCEL = f"人身损害赔偿统计数据_{START_YEAR}-{END_YEAR}.xlsx"
OUTPUT_JSON = "data_collection_status.json"
LOG_FILE = "data_collection.log"

# 可手工维护的公报URL覆盖项。
# 适合处理地方统计局搜索接口不可用、页面迁移或特殊栏目路径的情况。
# 示例：{("江苏省", 2024, "国民经济和社会发展统计公报"): "https://..."}
MANUAL_BULLETIN_URLS: Dict[Tuple[str, int, str], str] = {}

SUCCESS_STATUSES = {"已提取", "需人工复核"}
FAILED_STATUSES = {"未找到", "提取失败", "部分提取"}
COLLECTION_VERSION = "2026.2"

# ==================== 数据类定义 ====================

@dataclass
class DataRecord:
    """单条数据记录"""
    region: str
    year: int
    urban_disposable_income: Optional[float] = None  # 城镇居民人均可支配收入
    urban_consumption_expenditure: Optional[float] = None  # 城镇居民人均消费支出
    urban_non_private_wage: Optional[float] = None  # 城镇非私营单位年平均工资
    urban_private_wage: Optional[float] = None  # 城镇私营单位年平均工资
    service_industry_wage: Optional[float] = None  # 居民服务业年平均工资
    urban_on_duty_wage: Optional[float] = None  # 城镇单位在岗职工年平均工资
    source_url: Optional[str] = None
    publish_date: Optional[str] = None
    status: str = "待采集"  # 待采集、已提取、未找到、提取失败、需人工复核
    error_message: Optional[str] = None
    bulletin_type: Optional[str] = None  # 使用的公报类型
    evidence: Dict[str, Dict[str, Any]] = field(default_factory=dict)  # 字段级来源证据
    candidates: Dict[str, List[Dict[str, Any]]] = field(default_factory=dict)
    extraction_warnings: List[str] = field(default_factory=list)
    extraction_version: Optional[str] = None

    def to_dict(self) -> Dict:
        """转换为字典格式"""
        return asdict(self)

@dataclass
class CollectionStatus:
    """收集状态管理器，支持断点续爬"""
    records: List[DataRecord] = field(default_factory=list)
    completed_regions: Dict[str, List[int]] = field(default_factory=dict)
    start_time: Optional[str] = None
    end_time: Optional[str] = None

    def save(self, file_path: str):
        """原子保存状态，避免中断时留下半截JSON。"""
        data = {
            "records": [record.to_dict() for record in self.records],
            "completed_regions": self.completed_regions,
            "start_time": self.start_time,
            "end_time": self.end_time
        }
        target_path = os.path.abspath(file_path)
        os.makedirs(os.path.dirname(target_path), exist_ok=True)
        temp_path = target_path + ".tmp"
        with open(temp_path, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
            f.flush()
            os.fsync(f.fileno())
        os.replace(temp_path, target_path)

    def load(self, file_path: str):
        """从文件加载状态"""
        if os.path.exists(file_path):
            with open(file_path, 'r', encoding='utf-8') as f:
                data = json.load(f)

            self.records = []
            for record_dict in data.get("records", []):
                record = DataRecord(**record_dict)
                self.records.append(record)

            self.completed_regions = data.get("completed_regions", {})
            self.start_time = data.get("start_time")
            self.end_time = data.get("end_time")

    def is_completed(self, region: str, year: int, retry_failed: bool = True,
                     extraction_version: Optional[str] = None) -> bool:
        """检查某个地区某年份是否已完成。

        默认只跳过已成功提取或需要人工复核的数据；未找到、提取失败的数据
        会在下次断点续跑时重新尝试，避免临时网络问题被永久缓存。
        """
        record = self.get_record(region, year)
        if record:
            if extraction_version is not None and record.extraction_version != extraction_version:
                return False
            if retry_failed:
                return record.status in SUCCESS_STATUSES
            return True
        return year in self.completed_regions.get(region, [])

    def add_record(self, record: DataRecord):
        """添加或替换记录，并按状态维护完成索引"""
        replaced = False
        for index, existing in enumerate(self.records):
            if existing.region == record.region and existing.year == record.year:
                self.records[index] = record
                replaced = True
                break

        if not replaced:
            self.records.append(record)

        if record.region not in self.completed_regions:
            self.completed_regions[record.region] = []

        if record.status in SUCCESS_STATUSES and record.year not in self.completed_regions[record.region]:
            self.completed_regions[record.region].append(record.year)
        elif record.status in FAILED_STATUSES and record.year in self.completed_regions[record.region]:
            self.completed_regions[record.region].remove(record.year)

    def get_record(self, region: str, year: int) -> Optional[DataRecord]:
        """获取指定地区年份的记录"""
        for record in self.records:
            if record.region == region and record.year == year:
                return record
        return None

# ==================== 工具函数 ====================

def ensure_dependencies(packages: Optional[List[str]] = None):
    """确认运行采集器所需第三方依赖已安装。"""
    if packages is None:
        missing = sorted(set(MISSING_DEPENDENCIES))
    else:
        missing = sorted({name for name in MISSING_DEPENDENCIES if name in packages})
    if missing:
        raise RuntimeError(
            "缺少依赖："
            + ", ".join(missing)
            + "。请先运行 `python -m pip install -r requirements.txt`。"
        )

def setup_logging():
    """配置日志"""
    logging.basicConfig(
        level=logging.INFO,
        format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
        handlers=[
            logging.FileHandler(LOG_FILE, encoding='utf-8'),
            logging.StreamHandler()
        ]
    )
    return logging.getLogger(__name__)

def create_retry_session() -> requests.Session:
    """创建带重试机制的requests会话"""
    ensure_dependencies()
    session = requests.Session()

    # 设置默认请求头，模拟浏览器
    session.headers.update({
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
        'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,*/*;q=0.8',
        'Accept-Language': 'zh-CN,zh;q=0.9,en;q=0.8',
        'Accept-Encoding': 'gzip, deflate',
        'Connection': 'keep-alive',
        'Upgrade-Insecure-Requests': '1',
    })

    # 配置重试策略
    retry_strategy = Retry(
        total=MAX_RETRIES,
        backoff_factor=RETRY_BACKOFF_FACTOR,
        status_forcelist=[429, 500, 502, 503, 504],
        allowed_methods=["HEAD", "GET", "OPTIONS"]
    )
    adapter = HTTPAdapter(max_retries=retry_strategy)
    session.mount("http://", adapter)
    session.mount("https://", adapter)

    return session

def detect_encoding(content: bytes) -> str:
    """检测字节内容的编码"""
    if chardet is None:
        return 'utf-8'

    result = chardet.detect(content)
    return result.get('encoding') or 'utf-8'

def normalize_number(text: str) -> Optional[float]:
    """
    规范化数字字符串，处理千位分隔符和单位

    支持格式：
    - "42,568元" -> 42568.0
    - "3.5万元" -> 35000.0
    - "12 345" -> 12345.0
    """
    if not text:
        return None

    # 统一全角数字和常见中文标点，便于处理政府网站复制出的文本。
    text = text.translate(str.maketrans("０１２３４５６７８９，．－", "0123456789,.-"))
    text = re.sub(r'\s+', '', text)

    # 如果是"未发布"、"暂无"等，返回None
    if any(word in text.lower() for word in ['未发布', '尚未发布', '未公布', '暂无', '待公布', '—', '--', 'null']):
        return None

    # 处理"万元"或"万"单位
    is_ten_thousand = '万元' in text or '万' in text
    text = text.replace('万元', '').replace('万', '').replace('元', '')

    # 移除其他中文字符和符号
    text = re.sub(r'[^\d.,\-]', '', text)

    # 处理千位分隔符（逗号）
    text = text.replace(',', '')

    try:
        value = float(text)
        if is_ten_thousand:
            value *= 10000
        return value
    except ValueError:
        return None

NUMBER_WITH_UNIT_PATTERN = re.compile(
    r'(?P<number>[+-]?(?:\d{1,3}(?:[,，]\d{3})+|\d+)(?:[.．]\d+)?)\s*(?P<unit>万元|元|万)'
)
UNITLESS_NUMBER_PATTERN = re.compile(
    r'(?P<number>[+-]?(?:\d{1,3}(?:[,，]\d{3})+|\d+)(?:[.．]\d+)?)'
)
GROWTH_CONTEXT_WORDS = ("增长", "增加", "增收", "下降", "减少", "同比", "比上年", "增幅", "降幅")


def _looks_like_growth_number(window: str, start: int, end: int) -> bool:
    segment_start = max(
        window.rfind(separator, 0, start)
        for separator in ("，", ",", "。", "；", ";", "、", "：", ":", "\n")
    )
    context_before = window[max(0, segment_start + 1):start]
    if any(word in context_before for word in GROWTH_CONTEXT_WORDS):
        return True
    return window[end:end + 1] in {"%", "％"}


def _find_value_match_in_window(window: str) -> Optional[Tuple[float, int, int]]:
    for match in NUMBER_WITH_UNIT_PATTERN.finditer(window):
        if _looks_like_growth_number(window, match.start(), match.end()):
            continue
        value = normalize_number(match.group(0))
        if value is not None:
            return value, match.start(), match.end()

    # 部分统计表的表头已经标明“单位：元”，行内可能只剩数字。
    for match in UNITLESS_NUMBER_PATTERN.finditer(window):
        if _looks_like_growth_number(window, match.start(), match.end()):
            continue
        if window[match.end():match.end() + 1] == "年":
            continue
        value = normalize_number(match.group(0))
        if value is not None and value >= 3000:
            return value, match.start(), match.end()

    return None


def _find_value_in_window(window: str) -> Optional[float]:
    result = _find_value_match_in_window(window)
    if result is None:
        return None
    return result[0]


def _make_snippet(text: str, start: int, end: int, radius: int = 100) -> str:
    snippet = text[max(0, start - radius):min(len(text), end + radius)]
    return re.sub(r'\s+', ' ', snippet).strip()


def extract_number_with_evidence(text: str, keywords: List[str]) -> Tuple[Optional[float], Optional[Dict[str, Any]]]:
    """从文本中提取数字，并返回用于复核的命中片段。"""
    normalized_text = re.sub(r'\s+', ' ', text)
    for keyword in keywords:
        for match in re.finditer(re.escape(keyword), normalized_text):
            window_start = match.end()
            window = normalized_text[window_start:window_start + 120]
            value_match = _find_value_match_in_window(window)
            if value_match is None:
                continue

            value, value_start, value_end = value_match
            absolute_value_start = window_start + value_start
            absolute_value_end = window_start + value_end
            return value, {
                "keyword": keyword,
                "extract_method": "keyword_window",
                "snippet": _make_snippet(normalized_text, match.start(), absolute_value_end),
            }

    return None, None


def extract_number_from_text(text: str, keywords: List[str]) -> Optional[float]:
    """
    从文本中提取数字

    示例：从"城镇居民人均可支配收入 42,568元"中提取42568.0
    """
    value, _ = extract_number_with_evidence(text, keywords)
    return value

# ==================== 公报搜索与下载 ====================

class BulletinSearcher:
    """公报搜索器"""

    def __init__(self, session: requests.Session, logger: logging.Logger):
        self.session = session
        self.logger = logger

    def _build_url(self, base_url: str, path: str) -> str:
        """构建完整URL，处理相对路径"""
        return urllib.parse.urljoin(base_url, path)

    def _fetch_content(self, url: str, base_url: str) -> Tuple[Any, bytes]:
        return fetch_official_content(self.session, url, base_url)

    def _extract_links_from_content(self, content: str, page_url: str) -> List[Tuple[str, str, str]]:
        """从普通HTML和脚本CDATA中提取链接"""
        links = []
        seen = set()

        if BeautifulSoup is not None:
            soup = BeautifulSoup(content, 'html.parser')
            for link in soup.find_all('a', href=True):
                text = link.get_text(strip=True) or link.get('title', '').strip()
                href = urllib.parse.urljoin(page_url, link['href'])
                key = (text, href)
                if key not in seen:
                    links.append((text, href, link.parent.get_text(" ", strip=True) if link.parent else text))
                    seen.add(key)

        link_pattern = re.compile(r'<a\b(?P<attrs>[^>]*)>(?P<body>.*?)</a>', re.IGNORECASE | re.DOTALL)
        href_pattern = re.compile(r'href=["\'](?P<href>[^"\']+)["\']', re.IGNORECASE)
        title_pattern = re.compile(r'title=["\'](?P<title>[^"\']*)["\']', re.IGNORECASE)

        for match in link_pattern.finditer(content):
            attrs = match.group('attrs')
            href_match = href_pattern.search(attrs)
            if not href_match:
                continue
            title_match = title_pattern.search(attrs)
            title = title_match.group('title').strip() if title_match else ''
            body = re.sub(r'<[^>]+>', '', match.group('body')).strip()
            text = title or body
            href = urllib.parse.urljoin(page_url, href_match.group('href'))
            context = content[match.start():min(len(content), match.end() + 120)]
            key = (text, href)
            if text and key not in seen:
                links.append((text, href, context))
                seen.add(key)

        return links

    def _find_bulletin_link_in_content(
        self,
        content: str,
        page_url: str,
        year: int,
        bulletin_type: Dict,
        official_url: Optional[str] = None,
    ) -> Optional[Dict]:
        """在栏目页内容中寻找匹配公报链接"""
        for link_text, link_href, context in self._extract_links_from_content(content, page_url):
            if str(year) not in link_text:
                continue
            if not any(keyword in link_text for keyword in bulletin_type["keywords"]):
                continue
            if official_url and not is_allowed_official_url(link_href, official_url):
                continue

            return {
                "url": link_href,
                "type": bulletin_type["name"],
                "publish_date": self._extract_publish_date_from_text(context)
            }

        return None

    def _extract_candidate_column_urls(self, content: str, page_url: str, base_url: str) -> List[str]:
        """从页面中提取可能承载公报列表的栏目链接"""
        candidates = []
        for link_text, link_href, _ in self._extract_links_from_content(content, page_url):
            if any(keyword in link_text for keyword in ["统计公报", "年度统计公报", "数据发布", "平均工资", "居民收入"]):
                candidates.append(link_href)

        for match in re.finditer(r'location\.href\s*=\s*["\'](?P<href>[^"\']+)["\']', content, re.IGNORECASE):
            candidates.append(urllib.parse.urljoin(page_url, match.group('href')))

        for match in re.finditer(r'href=["\'](?P<href>[^"\']*dataproxy\.jsp[^"\']+)["\']', content, re.IGNORECASE):
            candidates.append(urllib.parse.urljoin(page_url, match.group('href')))

        common_column_paths = [
            "/col/col85275/index.html",
            "/col/col85666/index.html",
            "/col/col85764/index.html",
            "/tjsj/tjgb/index.html",
            "/tjgb/index.html",
            "/xxgk/tjgb/index.html",
            "/tjsj/index.html",
        ]
        candidates.extend(urllib.parse.urljoin(base_url, path) for path in common_column_paths)

        ordered = []
        seen = set()
        for url in candidates:
            if url not in seen and is_allowed_official_url(url, base_url):
                ordered.append(url)
                seen.add(url)
        return ordered

    def search_bulletin(self, region: str, year: int, bulletin_type: Dict) -> Optional[Dict]:
        """
        搜索指定地区、年份和类型的公报

        返回：{"url": 公报URL, "type": 公报类型, "publish_date": 发布日期}
        """
        base_url = REGION_STATS_URLS.get(region)
        if not base_url:
            self.logger.warning(f"地区 {region} 未配置统计局URL")
            return None

        manual_url = MANUAL_BULLETIN_URLS.get((region, year, bulletin_type["name"]))
        if manual_url:
            if is_allowed_official_url(manual_url, base_url):
                return {
                    "url": manual_url,
                    "type": bulletin_type["name"],
                    "publish_date": None,
                    "official_base_url": base_url,
                }
            self.logger.warning(f"忽略越出官方域名范围的手工公报URL：{manual_url}")

        # 尝试不同的搜索策略
        search_strategies = [
            self._search_by_column_pages,
            self._search_by_keywords,
            self._search_by_common_patterns,
            self._search_by_sitemap
        ]

        for strategy in search_strategies:
            result = strategy(region, year, bulletin_type, base_url)
            if result and is_allowed_official_url(result.get("url", ""), base_url):
                result["official_base_url"] = base_url
                return result

        return None

    def _search_by_column_pages(self, region: str, year: int, bulletin_type: Dict, base_url: str) -> Optional[Dict]:
        """通过首页和统计公报栏目页搜索"""
        queue = [base_url]
        seen = set()

        while queue and len(seen) < 12:
            page_url = queue.pop(0)
            if page_url in seen:
                continue
            seen.add(page_url)

            try:
                response, response_content = self._fetch_content(page_url, base_url)
                if response.status_code != 200:
                    continue

                encoding = detect_encoding(response_content)
                content = response_content.decode(encoding, errors='ignore')
                result = self._find_bulletin_link_in_content(
                    content,
                    response.url,
                    year,
                    bulletin_type,
                    official_url=base_url,
                )
                if result:
                    return result

                for candidate_url in self._extract_candidate_column_urls(content, response.url, base_url):
                    if candidate_url not in seen and candidate_url not in queue:
                        queue.append(candidate_url)
            except Exception as e:
                self.logger.debug(f"栏目页搜索失败 {page_url}: {str(e)}")
                continue

        return None

    def _search_by_keywords(self, region: str, year: int, bulletin_type: Dict, base_url: str) -> Optional[Dict]:
        """通过关键词搜索"""
        for keyword in bulletin_type["keywords"]:
            # 构造搜索URL（不同网站搜索接口不同，这里使用通用模式）
            encoded_keyword = urllib.parse.quote(keyword)
            search_urls = [
                self._build_url(base_url, f"/search?q={year}{encoded_keyword}"),
                self._build_url(base_url, f"/search?keyword={year}{encoded_keyword}"),
                self._build_url(base_url, f"/search?wd={year}{encoded_keyword}"),
            ]

            for search_url in search_urls:
                try:
                    response, response_content = self._fetch_content(search_url, base_url)
                    if response.status_code == 200:
                        soup = BeautifulSoup(
                            response_content,
                            'html.parser',
                            from_encoding=detect_encoding(response_content),
                        )

                        # 查找公报链接（通常包含"公报"、"公告"等字样）
                        links = soup.find_all('a', href=True)
                        for link in links:
                            link_text = link.get_text(strip=True)
                            link_href = link['href']

                            # 检查链接是否匹配公报特征
                            if (str(year) in link_text and
                                any(kw in link_text for kw in bulletin_type["keywords"]) and
                                ('公报' in link_text or '公告' in link_text or '统计' in link_text)):

                                # 处理相对URL，使用_build_url确保正确构建完整URL
                                link_href = self._build_url(response.url, link_href)
                                if not is_allowed_official_url(link_href, base_url):
                                    continue

                                # 尝试获取发布日期
                                publish_date = self._extract_publish_date(link.parent)

                                return {
                                    "url": link_href,
                                    "type": bulletin_type["name"],
                                    "publish_date": publish_date
                                }
                except Exception as e:
                    self.logger.debug(f"搜索失败 {search_url}: {str(e)}")
                    continue

        return None

    def _search_by_common_patterns(self, region: str, year: int, bulletin_type: Dict, base_url: str) -> Optional[Dict]:
        """通过常见URL模式搜索"""
        # 常见公报URL模式
        common_patterns = [
            f"{base_url}/tjsj/tjgb/{year}/index.html",
            f"{base_url}/tjsj/tjgb/{year}/",
            f"{base_url}/tjgb/{year}/",
            f"{base_url}/xxgk/tjgb/{year}/",
            f"{base_url}/tjj/tjgb/{year}/",
        ]

        for pattern in common_patterns:
            try:
                response, response_content = self._fetch_content(pattern, base_url)
                if response.status_code == 200:
                    soup = BeautifulSoup(
                        response_content,
                        'html.parser',
                        from_encoding=detect_encoding(response_content),
                    )

                    # 在页面中查找公报链接
                    for keyword in bulletin_type["keywords"]:
                        links = soup.find_all('a', href=True, string=re.compile(rf'.*{year}.*{re.escape(keyword)}.*'))
                        if links:
                            link = links[0]
                            link_href = link['href']

                            # 处理相对URL，使用当前页面URL作为基础
                            link_href = urllib.parse.urljoin(response.url, link_href)
                            if not is_allowed_official_url(link_href, base_url):
                                continue

                            publish_date = self._extract_publish_date(link.parent)

                            return {
                                "url": link_href,
                                "type": bulletin_type["name"],
                                "publish_date": publish_date
                            }
            except Exception as e:
                self.logger.debug(f"模式匹配失败 {pattern}: {str(e)}")
                continue

        return None

    def _search_by_sitemap(self, region: str, year: int, bulletin_type: Dict, base_url: str) -> Optional[Dict]:
        """通过网站地图搜索"""
        sitemap_urls = [
            f"{base_url}/sitemap.xml",
            f"{base_url}/sitemap.txt",
            f"{base_url}/map.html",
        ]

        for sitemap_url in sitemap_urls:
            try:
                response, response_content = self._fetch_content(sitemap_url, base_url)
                if response.status_code == 200:
                    # 简单查找包含年份和关键词的URL
                    content = response_content.decode(detect_encoding(response_content), errors="ignore")
                    for keyword in bulletin_type["keywords"]:
                        pattern = rf'([^"\']*{year}[^"\']*{re.escape(keyword)}[^"\']*\.(?:html|pdf|htm)[^"\']*)'
                        matches = re.findall(pattern, content, re.IGNORECASE)
                        if matches:
                            url = matches[0]
                            # 处理相对URL
                            url = urllib.parse.urljoin(response.url, url)
                            if not is_allowed_official_url(url, base_url):
                                continue

                            return {
                                "url": url,
                                "type": bulletin_type["name"],
                                "publish_date": None
                            }
            except Exception as e:
                self.logger.debug(f"网站地图搜索失败 {sitemap_url}: {str(e)}")
                continue

        return None

    def _extract_publish_date_from_text(self, text: str) -> Optional[str]:
        """从文本中提取发布日期"""
        # 常见日期模式
        date_patterns = [
            r'(\d{4})年(\d{1,2})月(\d{1,2})日',
            r'(\d{4})-(\d{1,2})-(\d{1,2})',
            r'(\d{4})/(\d{1,2})/(\d{1,2})',
        ]

        for pattern in date_patterns:
            match = re.search(pattern, text)
            if match:
                year, month, day = match.groups()
                return f"{year}-{int(month):02d}-{int(day):02d}"

        return None

    def _extract_publish_date(self, element) -> Optional[str]:
        """从HTML元素中提取发布日期"""
        if not element:
            return None

        return self._extract_publish_date_from_text(element.get_text(strip=True))

# ==================== 数据提取器 ====================

class DataExtractor:
    """数据提取器"""

    def __init__(self, session: requests.Session, logger: logging.Logger, archive=None):
        self.session = session
        self.logger = logger or logging.getLogger(__name__)
        self.archive = archive or ArchiveStore()
        self._document_cache = OrderedDict()

    def extract_data(self, bulletin_info: Dict) -> Dict[str, Any]:
        """Extract candidates against explicit dimensions and archive original bytes."""
        from personal_injury.collection.extraction import extract_document

        url = bulletin_info.get("url")
        bulletin_type = bulletin_info.get("type")
        publish_date = bulletin_info.get("publish_date")
        official_base_url = bulletin_info.get("official_base_url")
        empty = {"data": {}, "evidence": {}, "candidates": {}, "warnings": [],
                 "source_url": url, "publish_date": publish_date, "bulletin_type": bulletin_type}

        try:
            if not url or not official_base_url:
                raise ValueError("公报信息缺少official_base_url，拒绝执行未限定域名的抓取")
            region, year = bulletin_info.get("region"), bulletin_info.get("year")
            if not region or not isinstance(year, int) or isinstance(year, bool):
                raise ValueError("抽取必须指定地区和统计年度，不能由发布日期推断")
            cache_key = (url, official_base_url)
            if cache_key in self._document_cache:
                final_url, content_type, response_content, retrieved_at = self._document_cache[cache_key]
                self._document_cache.move_to_end(cache_key)
            else:
                response, response_content = fetch_official_content(self.session, url, official_base_url)
                response.raise_for_status()
                final_url = response.url
                content_type = response.headers.get("Content-Type", "").lower()
                retrieved_at = datetime.now().isoformat(timespec='seconds')
                self._document_cache[cache_key] = (final_url, content_type, response_content, retrieved_at)
                if len(self._document_cache) > 4:
                    self._document_cache.popitem(last=False)
            source_title = bulletin_info.get("title")
            if "html" in content_type or response_content.lstrip().startswith((b"<", b"<!")):
                soup = BeautifulSoup(response_content, 'html.parser', from_encoding=detect_encoding(response_content))
                # Search result titles can supply a missing HTML title, but never override one.
                heading = soup.find("h1") or soup.title
                if heading:
                    source_title = heading.get_text(" ", strip=True)
            archived = self.archive.save(response_content, source_url=final_url, content_type=content_type,
                                         metadata={"retrieved_at": retrieved_at, "title": source_title})
            result = extract_document(response_content, content_type, region=region, year=year,
                                      title=source_title, source_url=final_url)
            source_metadata = {
                **archived,
                "source_url": final_url,
                "source_title": source_title,
                "publish_date": publish_date,
                "bulletin_type": bulletin_type,
                "retrieved_at": retrieved_at,
                "source_domain_verified": is_allowed_official_url(final_url, official_base_url),
                "transport_secure": urllib.parse.urlsplit(final_url).scheme == "https",
            }
            result["evidence"] = {name: {**item, **source_metadata}
                                  for name, item in result.get("evidence", {}).items()}
            result["candidates"] = {name: [{**item, **source_metadata} for item in items]
                                    for name, items in result.get("candidates", {}).items()}
            result.update({key: source_metadata[key] for key in
                           ("source_url", "source_title", "publish_date", "bulletin_type")})
            attachments = []
            if 'soup' in locals():
                for link in soup.find_all("a", href=True):
                    attachment = urllib.parse.urljoin(final_url, link["href"])
                    if (urllib.parse.urlsplit(attachment).path.lower().endswith((".pdf", ".xlsx"))
                            and is_allowed_official_url(attachment, official_base_url)
                            and attachment not in [item["url"] for item in attachments]):
                        attachments.append({**bulletin_info, "url": attachment, "title": source_title})
                    if len(attachments) >= 4:
                        break
            result["attachments"] = attachments
            return result
        except Exception as e:
            self.logger.error(f"提取数据失败 {url}: {str(e)}")
            return {**empty, "warnings": [str(e)]}

# ==================== 数据校验器 ====================

class DataValidator:
    """数据校验器"""

    def __init__(self, logger: logging.Logger):
        self.logger = logger

    def _check_value_ranges(self, record: DataRecord) -> List[str]:
        """检查数值是否落在保守合理范围内，防止抓到年份、增长率等错误数字"""
        warnings = []
        for item_name, attr_name in DATA_FIELD_ATTRS.items():
            value = getattr(record, attr_name)
            if value is None:
                continue

            min_value, max_value = VALUE_RANGE_RULES[item_name]
            if value < min_value or value > max_value:
                warnings.append(f"{item_name}={value:g}超出合理范围{min_value:g}-{max_value:g}")
        return warnings

    def validate_record(self, record: DataRecord, all_records: List[DataRecord],
                        required_items: Optional[List[str]] = None) -> DataRecord:
        """
        验证单条记录的合理性

        检查项目：
        1. 关键数据项是否缺失
        2. 同比变化是否合理（±30%以内）
        3. 数据间逻辑关系是否合理
        """
        missing_required_items = [
            item_name
            for item_name in (required_items if required_items is not None else REQUIRED_DATA_ITEMS)
            if getattr(record, DATA_FIELD_ATTRS[item_name]) is None
        ]
        missing_optional_items = [
            item_name
            for item_name in (OPTIONAL_LEGACY_DATA_ITEMS if required_items is None else [])
            if getattr(record, DATA_FIELD_ATTRS[item_name]) is None
        ]
        present_items = [
            item_name
            for item_name, attr_name in DATA_FIELD_ATTRS.items()
            if getattr(record, attr_name) is not None
        ]
        range_warnings = self._check_value_ranges(record)
        unverified_source_items = []
        insecure_transport_items = []
        conflicting_items = []
        dimension_warnings = []
        if record.evidence:
            unverified_source_items = [
                item_name
                for item_name in present_items
                if record.evidence.get(item_name, {}).get("source_domain_verified") is not True
            ]
            insecure_transport_items = [
                item_name
                for item_name in present_items
                if record.evidence.get(item_name, {}).get("transport_secure") is False
            ]
            conflicting_items = [
                item_name
                for item_name in record.evidence
                if record.evidence.get(item_name, {}).get("conflicts")
            ]
            if record.extraction_version:
                dimension_warnings = [item_name for item_name in present_items if (
                    record.evidence.get(item_name, {}).get("validation_status") != "verified"
                    or record.evidence.get(item_name, {}).get("region") != record.region
                    or record.evidence.get(item_name, {}).get("statistical_year") != record.year
                    or record.evidence.get(item_name, {}).get("unit") != "元"
                )]

        notes = []
        if missing_required_items:
            notes.append("缺失核心数据项：" + "、".join(missing_required_items))
        if missing_optional_items:
            notes.append("缺失可选历史口径：" + "、".join(missing_optional_items))
        if range_warnings:
            notes.append("异常数值：" + "；".join(range_warnings))
        if unverified_source_items:
            notes.append("来源域名未通过校验：" + "、".join(unverified_source_items))
        if insecure_transport_items:
            notes.append("来源未使用HTTPS传输：" + "、".join(insecure_transport_items))
        if conflicting_items:
            notes.append("同一字段存在冲突候选值：" + "、".join(conflicting_items))
        if dimension_warnings:
            notes.append("抽取口径或维度未经确认：" + "、".join(dimension_warnings))

        if not present_items:
            record.status = "需人工复核" if record.candidates else "提取失败"
            record.error_message = "；".join(["没有唯一确认的数值，候选见复核清单" if record.candidates else "所有数据项缺失", *notes])
            return record

        if missing_required_items:
            record.status = "部分提取"
            record.error_message = "；".join(notes)
            return record

        if range_warnings or unverified_source_items or insecure_transport_items or conflicting_items or dimension_warnings:
            record.status = "需人工复核"
            record.error_message = "；".join(notes)
            return record

        # 检查同比变化
        needs_review = self._check_year_over_year(record, all_records)

        if needs_review:
            record.status = "需人工复核"
            notes.append("同比变化超过30%")
            if notes:
                record.error_message = "；".join(notes)
        elif notes:
            record.status = "已提取"
            record.error_message = "；".join(notes)

        return record

    def _check_year_over_year(self, record: DataRecord, all_records: List[DataRecord]) -> bool:
        """检查同比增长率，超过±30%返回True"""
        # 查找同地区上一年的数据
        prev_year_record = None
        for r in all_records:
            if r.region == record.region and r.year == record.year - 1:
                prev_year_record = r
                break

        if not prev_year_record:
            return False

        # 检查各项数据的同比增长率
        data_items = [
            "urban_disposable_income",
            "urban_consumption_expenditure",
            "urban_non_private_wage",
            "urban_private_wage",
            "service_industry_wage",
            "urban_on_duty_wage"
        ]

        for item in data_items:
            current_val = getattr(record, item)
            prev_val = getattr(prev_year_record, item)

            if current_val is not None and prev_val is not None and prev_val != 0:
                growth_rate = (current_val - prev_val) / prev_val
                if abs(growth_rate) > 0.3:  # ±30%
                    (self.logger or logging.getLogger(__name__)).warning(
                        f"数据异常增长：{record.region} {record.year} {item} "
                        f"增长率{growth_rate:.1%}"
                    )
                    return True

        return False

# ==================== 主收集器 ====================

class DataCollector:
    """主数据收集器"""

    def __init__(self, *, source: str = "auto", archive_dir=None, status_file: str = OUTPUT_JSON):
        ensure_dependencies()
        if source not in {"auto", "nbs", "bulletin"}:
            raise ValueError("source必须是auto、nbs或bulletin")
        self.logger = setup_logging()
        self.session = create_retry_session()
        self.status = CollectionStatus()
        self.source = source
        self.status_file = status_file
        self.archive = ArchiveStore(archive_dir)

        # 初始化组件
        self.searcher = BulletinSearcher(self.session, self.logger)
        self.extractor = DataExtractor(self.session, self.logger, archive=self.archive)
        self.validator = DataValidator(self.logger)
        from personal_injury.collection.nbs import NbsSource
        # Isolate NBS's short, bounded requests from the bulletin crawler's retries.
        self.nbs = NbsSource(requests.Session(), fetch_content=fetch_official_content,
                             archive=self.archive, logger=self.logger)

    def _is_official_source(self, region: str, source_url: Optional[str]) -> bool:
        """检查来源域名是否匹配该地区配置的统计局官网域名。"""
        if not source_url:
            return False

        official_url = REGION_STATS_URLS.get(region)
        if not official_url:
            return False

        return (is_allowed_official_url(source_url, official_url)
                or is_allowed_official_url(source_url, "https://data.stats.gov.cn"))

    def _merge_result(self, record: DataRecord, result: Dict[str, Any]) -> None:
        """Preserve all candidates and block conflicting values rather than choosing one."""
        record.extraction_warnings.extend(result.get("warnings", []))
        for item_name, candidates in result.get("candidates", {}).items():
            record.candidates.setdefault(item_name, []).extend(candidates)
        for item_name, item_evidence in result.get("evidence", {}).items():
            if item_name not in record.evidence:
                record.evidence[item_name] = dict(item_evidence)
            elif item_evidence.get("conflicts"):
                record.evidence[item_name].setdefault("conflicts", []).extend(
                    item_evidence["conflicts"]
                )
            if item_evidence.get("conflicts"):
                record.evidence[item_name]["validation_status"] = "pending"
                if item_name in DATA_FIELD_ATTRS:
                    setattr(record, DATA_FIELD_ATTRS[item_name], None)
        for item_name, value in result.get("data", {}).items():
            if value is None or item_name not in DATA_FIELD_ATTRS:
                continue
            evidence = dict(result.get("evidence", {}).get(item_name, {}))
            if (evidence.get("validation_status") != "verified"
                    or evidence.get("region") != record.region
                    or evidence.get("statistical_year") != record.year
                    or evidence.get("unit") != "元"):
                record.extraction_warnings.append(f"{item_name}未通过抽取维度校验，未写入数值")
                continue
            evidence["source_domain_verified"] = self._is_official_source(record.region, evidence.get("source_url"))
            attr = DATA_FIELD_ATTRS[item_name]
            previous = getattr(record, attr)
            previous_evidence = record.evidence.get(item_name, {})
            if previous_evidence.get("conflicts"):
                previous_evidence["conflicts"].append({**evidence, "value": value})
                continue
            if previous is not None and Decimal(str(previous)) != Decimal(str(value)):
                previous_evidence["conflicts"] = [{**previous_evidence, "value": previous}, {**evidence, "value": value}]
                previous_evidence["validation_status"] = "pending"
                previous_evidence.setdefault("review_reasons", []).append("官方来源数值不一致，需核对更正与口径")
                setattr(record, attr, None)
            elif previous is None:
                setattr(record, attr, value)
                record.evidence[item_name] = evidence
            else:
                previous_evidence.setdefault("corroborating_sources", []).append(evidence)

    def collect_record(self, region: str, year: int) -> DataRecord:
        record = DataRecord(region=region, year=year, extraction_version=f"{COLLECTION_VERSION}:{self.source}")
        source_urls, publish_dates, bulletin_types = [], [], []

        def merge(result):
            self._merge_result(record, result)
            for key, destination in (("source_url", source_urls), ("publish_date", publish_dates),
                                     ("bulletin_type", bulletin_types)):
                if result.get(key):
                    destination.append(result[key])

        if self.source in {"auto", "nbs"}:
            merge(self.nbs.collect(region, year))
        if self.source in {"auto", "bulletin"}:
            seen_urls = set()
            for bulletin_type in BULLETIN_TYPES:
                bulletin = self.searcher.search_bulletin(region, year, bulletin_type)
                if not bulletin or bulletin.get("url") in seen_urls:
                    continue
                seen_urls.add(bulletin["url"])
                bulletin.update(region=region, year=year)
                result = self.extractor.extract_data(bulletin)
                merge(result)
                for attachment in result.get("attachments", []):
                    if attachment["url"] not in seen_urls:
                        seen_urls.add(attachment["url"])
                        merge(self.extractor.extract_data(attachment))
        record.source_url = "; ".join(dict.fromkeys(source_urls)) or None
        record.publish_date = "; ".join(dict.fromkeys(publish_dates)) or None
        record.bulletin_type = "; ".join(dict.fromkeys(bulletin_types)) or None
        record.extraction_warnings = list(dict.fromkeys(record.extraction_warnings))
        record.status = "已提取"
        required = ["城镇居民人均可支配收入", "城镇居民人均消费支出"] if self.source == "nbs" else None
        record = self.validator.validate_record(record, self.status.records, required_items=required)
        if not source_urls and not record.candidates:
            record.status = "未找到"
        if record.extraction_warnings:
            record.error_message = "；".join(filter(None, [record.error_message, *record.extraction_warnings]))
        return record

    def run(self, start_year: int = START_YEAR, end_year: int = END_YEAR,
             regions: List[str] = None, output_file: str = None, resume: bool = True,
             retry_failed: bool = True):
        """运行数据收集

        Args:
            start_year: 起始年份
            end_year: 结束年份
            regions: 地区列表，默认为所有地区
            output_file: 输出Excel文件名，默认为配置中的OUTPUT_EXCEL
            resume: 是否从上次进度恢复，默认为True
            retry_failed: 断点续跑时是否重试失败记录，默认为True
        """
        if isinstance(start_year, bool) or isinstance(end_year, bool):
            raise ValueError("起止年份必须是整数")
        if not isinstance(start_year, int) or not isinstance(end_year, int):
            raise ValueError("起止年份必须是整数")
        if start_year > end_year:
            raise ValueError("起始年份不能大于结束年份")

        if regions is None:
            regions = list(REGIONS)
        else:
            regions = list(regions)
        unsupported_regions = [region for region in regions if region not in REGION_STATS_URLS]
        if unsupported_regions:
            raise ValueError("未配置统计局官网的地区：" + "、".join(unsupported_regions))

        if output_file is None:
            output_file = OUTPUT_EXCEL

        self.logger.info("=" * 60)
        self.logger.info("开始收集人身损害赔偿统计数据")
        self.logger.info(f"地区数量：{len(regions)}")
        self.logger.info(f"年份范围：{start_year}-{end_year}")
        self.logger.info(f"输出文件：{output_file}")
        self.logger.info("=" * 60)

        # 加载之前的进度（断点续爬）
        if resume and os.path.exists(self.status_file):
            self.status.load(self.status_file)
            self.logger.info(f"加载了之前的进度，已完成{len(self.status.records)}条记录")
        elif not resume:
            self.logger.info("不加载之前进度，重新开始收集")
            self.status = CollectionStatus()  # 重置状态

        self.status.start_time = datetime.now().isoformat()

        # 遍历所有地区和年份
        total_tasks = len(regions) * (end_year - start_year + 1)
        completed_tasks = len(self.status.records)

        self.logger.info(f"总任务数：{total_tasks}，已完成：{completed_tasks}")

        for region in regions:
            self.logger.info(f"处理地区：{region}")

            for year in range(start_year, end_year + 1):
                # 检查是否已完成
                if self.status.is_completed(region, year, retry_failed=retry_failed,
                                            extraction_version=f"{COLLECTION_VERSION}:{self.source}"):
                    self.logger.debug(f"跳过已完成：{region} {year}")
                    continue

                self.logger.info(f"采集 {region} {year}年数据...")

                # 创建数据记录
                record = DataRecord(region=region, year=year)

                try:
                    record = self.collect_record(region, year)

                    # 4. 保存记录
                    self.status.add_record(record)

                    self.logger.info(f"处理完成：{region} {year}年，状态：{record.status}")

                except Exception as e:
                    self.logger.error(f"采集失败 {region} {year}年: {str(e)}")
                    record.status = "提取失败"
                    record.error_message = str(e)
                    self.status.add_record(record)
                finally:
                    # 即使本轮提前continue，也要定期保存并保持礼貌性延迟。
                    if len(self.status.records) % 10 == 0:
                        self.status.save(self.status_file)
                    time.sleep(1)

        # 最终保存
        self.status.end_time = datetime.now().isoformat()
        self.status.save(self.status_file)

        self.logger.info("数据收集完成")

        # 导出Excel
        self.export_to_excel(output_file)

        self.logger.info(f"结果已导出到：{output_file}")
        self.logger.info(f"详细日志：{LOG_FILE}")

    def _build_source_rows(self) -> List[Dict[str, Any]]:
        """生成字段级溯源明细行"""
        rows = []
        for record in self.status.records:
            for item_name, attr_name in DATA_FIELD_ATTRS.items():
                value = getattr(record, attr_name)
                evidence = record.evidence.get(item_name, {}) if record.evidence else {}
                requirement = "核心必需" if item_name in REQUIRED_DATA_ITEMS else "可选历史口径"
                rows.append({
                    "地区名称": record.region,
                    "统计年度": record.year,
                    "数据项": item_name,
                    "字段要求": requirement,
                    "数值": value,
                    "来源链接": evidence.get("source_url") or record.source_url,
                    "来源标题": evidence.get("source_title"),
                    "发布时间": evidence.get("publish_date") or record.publish_date,
                    "公报类型": evidence.get("bulletin_type") or record.bulletin_type,
                    "采集时间": evidence.get("retrieved_at"),
                    "内容SHA256": evidence.get("content_sha256"),
                    "官方域名校验": evidence.get("source_domain_verified"),
                    "HTTPS传输": evidence.get("transport_secure"),
                    "提取方法": evidence.get("extract_method"),
                    "命中关键词": evidence.get("keyword"),
                    "提取片段": evidence.get("snippet"),
                    "原始数值": evidence.get("raw_value"),
                    "原始单位": evidence.get("raw_unit"),
                    "归一单位": evidence.get("unit"),
                    "单位换算": json.dumps(evidence.get("conversion"), ensure_ascii=False),
                    "指标口径": json.dumps(evidence.get("scope"), ensure_ascii=False),
                    "原文位置": json.dumps(evidence.get("locator"), ensure_ascii=False),
                    "原文归档": evidence.get("archive_path"),
                    "抽取版本": evidence.get("extractor_version"),
                    "抽取校验": evidence.get("validation_status"),
                    "复核原因": "；".join(evidence.get("review_reasons", [])),
                    "冲突候选": json.dumps(evidence.get("conflicts", []), ensure_ascii=False),
                    "数据状态": record.status,
                    "错误信息": record.error_message,
                })
        return rows

    def _build_candidate_rows(self) -> List[Dict[str, Any]]:
        rows = []
        for record in self.status.records:
            for item_name, candidates in record.candidates.items():
                for candidate in candidates:
                    rows.append({
                        "目标地区": record.region, "目标统计年度": record.year, "数据项": item_name,
                        "候选数值": candidate.get("value"), "原始数值": candidate.get("raw_value"),
                        "原始单位": candidate.get("raw_unit"), "候选地区": candidate.get("region"),
                        "候选年度": candidate.get("statistical_year"),
                        "指标口径": json.dumps(candidate.get("scope"), ensure_ascii=False),
                        "抽取校验": candidate.get("validation_status"),
                        "复核原因": "；".join(candidate.get("review_reasons", [])),
                        "来源链接": candidate.get("source_url"), "原文归档": candidate.get("archive_path"),
                        "原文位置": json.dumps(candidate.get("locator"), ensure_ascii=False),
                        "原文片段": candidate.get("snippet"),
                    })
        return rows

    def _format_worksheet(self, worksheet):
        """设置工作表列宽和标题样式"""
        for column in worksheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except Exception:
                    pass
            adjusted_width = min(max_length + 2, 60)
            worksheet.column_dimensions[column_letter].width = adjusted_width

        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

    def export_to_excel(self, output_file: str = None):
        """导出数据到Excel

        Args:
            output_file: 输出Excel文件名，默认为配置中的OUTPUT_EXCEL
        """
        if output_file is None:
            output_file = OUTPUT_EXCEL
        output_path = os.path.abspath(output_file)
        os.makedirs(os.path.dirname(output_path), exist_ok=True)

        # 准备数据
        rows = []
        for record in self.status.records:
            row = {
                "地区名称": record.region,
                "统计年度": record.year,
                "城镇居民人均可支配收入(元)": record.urban_disposable_income,
                "城镇居民人均消费支出(元)": record.urban_consumption_expenditure,
                "城镇非私营单位年平均工资(元)": record.urban_non_private_wage,
                "城镇私营单位年平均工资(元)": record.urban_private_wage,
                "居民服务等行业非私营就业人员年平均工资(元)": record.service_industry_wage,
                "城镇单位在岗职工年平均工资(元)": record.urban_on_duty_wage,
                "数据来源链接": record.source_url,
                "发布时间": record.publish_date,
                "数据状态": record.status,
                "公报类型": record.bulletin_type,
                "错误信息": record.error_message
            }
            rows.append(row)

        # 创建DataFrame
        df = pd.DataFrame(spreadsheet_safe_rows(rows))
        source_df = pd.DataFrame(spreadsheet_safe_rows(self._build_source_rows()))
        candidate_df = pd.DataFrame(spreadsheet_safe_rows(self._build_candidate_rows()))

        # 调整列顺序
        column_order = [
            "地区名称", "统计年度",
            "城镇居民人均可支配收入(元)", "城镇居民人均消费支出(元)",
            "城镇非私营单位年平均工资(元)", "城镇私营单位年平均工资(元)",
            "居民服务等行业非私营就业人员年平均工资(元)", "城镇单位在岗职工年平均工资(元)",
            "数据来源链接", "发布时间", "数据状态", "公报类型", "错误信息"
        ]
        df = df.reindex(columns=column_order)

        source_column_order = [
            "地区名称", "统计年度", "数据项", "字段要求", "数值",
            "来源链接", "来源标题", "发布时间", "公报类型",
            "采集时间", "内容SHA256", "官方域名校验",
            "HTTPS传输",
            "提取方法", "命中关键词", "提取片段", "冲突候选",
            "原始数值", "原始单位", "归一单位", "单位换算", "指标口径", "原文位置",
            "原文归档", "抽取版本", "抽取校验", "复核原因",
            "数据状态", "错误信息"
        ]
        source_df = source_df.reindex(columns=source_column_order)

        # 导出到Excel
        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='统计数据', index=False)
            source_df.to_excel(writer, sheet_name='字段溯源', index=False)
            candidate_df.to_excel(writer, sheet_name='候选复核', index=False)

            self._format_worksheet(writer.sheets['统计数据'])
            self._format_worksheet(writer.sheets['字段溯源'])
            self._format_worksheet(writer.sheets['候选复核'])

        # 同时保存为CSV以便查看
        csv_output = os.path.splitext(output_path)[0] + '.csv'
        source_csv_output = os.path.splitext(output_path)[0] + '_sources.csv'
        df.to_csv(csv_output, index=False, encoding='utf-8-sig')
        source_df.to_csv(source_csv_output, index=False, encoding='utf-8-sig')
        candidate_df.to_csv(os.path.splitext(output_path)[0] + '_candidates.csv', index=False, encoding='utf-8-sig')

# ==================== 主程序入口 ====================

if __name__ == "__main__":
    # 解析命令行参数
    parser = argparse.ArgumentParser(description='全国人身损害赔偿统计数据收集工具')
    parser.add_argument('--start-year', type=int, default=START_YEAR,
                       help=f'起始年份（默认：{START_YEAR}）')
    parser.add_argument('--end-year', type=int, default=END_YEAR,
                       help=f'结束年份（默认：{END_YEAR}）')
    parser.add_argument('--last-years', type=int, default=None,
                       help='收集最近N个完整统计年度，会根据结束年份自动计算起始年份')
    parser.add_argument('--region', type=str, default=None,
                       help='指定地区（默认：所有地区）')
    parser.add_argument('--no-resume', action='store_true',
                       help='不从上次进度恢复（默认：自动恢复）')
    parser.add_argument('--no-retry-failed', action='store_true',
                       help='断点续跑时不重试未找到或提取失败的记录（默认：重试失败记录）')
    parser.add_argument('--output', type=str, default=OUTPUT_EXCEL,
                       help=f'输出Excel文件名（默认：{OUTPUT_EXCEL}）')
    parser.add_argument('--source', choices=['auto', 'nbs', 'bulletin'], default='auto',
                       help='auto先查统计局结构化数据再核对地方公报；nbs仅查省级城镇收入和消费')
    parser.add_argument('--archive-dir', default=None, help='原文归档目录，默认data/source_archive或STATISTICAL_ARCHIVE_DIR')
    parser.add_argument('--status-file', default=OUTPUT_JSON, help='采集进度JSON路径')

    args = parser.parse_args()

    if args.last_years is not None:
        if args.last_years <= 0:
            print("错误：--last-years 必须大于0")
            exit(1)
        args.start_year = args.end_year - args.last_years + 1

    # 验证参数
    if args.start_year > args.end_year:
        print("错误：起始年份不能大于结束年份")
        exit(1)

    if args.region and args.region not in REGIONS:
        print(f"错误：地区 '{args.region}' 不在支持列表中")
        print(f"支持的地区：{', '.join(REGIONS[:5])}...等共{len(REGIONS)}个地区")
        exit(1)

    # 如果指定了地区，只处理该地区
    target_regions = [args.region] if args.region else REGIONS

    # 创建收集器
    try:
        collector = DataCollector(source=args.source, archive_dir=args.archive_dir, status_file=args.status_file)
    except RuntimeError as e:
        print(f"错误：{e}")
        exit(1)

    try:
        # 运行收集器，传递年份参数
        collector.run(start_year=args.start_year, end_year=args.end_year,
                       regions=target_regions, output_file=args.output,
                       resume=not args.no_resume,
                       retry_failed=not args.no_retry_failed)

        print("\n" + "=" * 60)
        print("数据收集完成！")
        print(f"结果文件：{args.output}")
        print(f"日志文件：{LOG_FILE}")
        print(f"进度文件：{collector.status_file}")
        print("=" * 60)

    except KeyboardInterrupt:
        print("\n用户中断，保存进度...")
        collector.status.save(collector.status_file)
        print(f"进度已保存到 {OUTPUT_JSON}")
    except Exception as e:
        print(f"程序运行出错：{str(e)}")
        import traceback
        traceback.print_exc()
