from __future__ import annotations

import copy
import json
import logging
from pathlib import Path
from unittest.mock import Mock

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine

from data_collector import (
    COLLECTION_VERSION, CollectionStatus, DataCollector, DataExtractor, DataRecord,
    DataValidator, fetch_official_content,
)
from personal_injury.collection.source_archive import ArchiveStore
from personal_injury.infrastructure.collector_import import import_collector_status
from personal_injury.infrastructure.database import Base, StatisticalStandardModel, create_session_factory


INCOME = "城镇居民人均可支配收入"
CONSUMPTION = "城镇居民人均消费支出"
URL = "https://tj.jiangsu.gov.cn/test.html"


def result(value=60000, *, year=2024, region="江苏省"):
    evidence = {
        "validation_status": "verified", "region": region, "statistical_year": year,
        "unit": "元", "raw_unit": "元", "raw_value": str(value), "value": value,
        "scope": {"population": "城镇居民", "metric": "人均可支配收入"},
        "source_url": URL, "source_domain_verified": True, "transport_secure": True,
        "snippet": f"城镇居民人均可支配收入{value}元", "locator": {"paragraph": 1},
        "extractor_version": "test/1", "review_reasons": [],
    }
    return {"data": {INCOME: value}, "evidence": {INCOME: evidence},
            "candidates": {INCOME: [copy.deepcopy(evidence)]}, "source_url": URL,
            "bulletin_type": "统计公报", "warnings": []}


def collector():
    instance = DataCollector.__new__(DataCollector)
    instance.source = "nbs"
    instance.status = CollectionStatus()
    instance.validator = DataValidator(logging.getLogger(__name__))
    return instance


class Response:
    status_code = 200
    url = URL
    headers = {"Content-Type": "text/html; charset=utf-8"}

    def __init__(self, content):
        self.content = content

    def raise_for_status(self):
        pass

    def iter_content(self, chunk_size):
        yield self.content

    def close(self):
        pass


def test_real_html_pipeline_keeps_correct_year_and_archives_bytes(tmp_path):
    original = f"""<html><h1>江苏省2024年统计公报</h1><table><caption>单位：万元</caption>
    <tr><th>指标</th><th>2023年</th><th>2024年</th></tr>
    <tr><td>{INCOME}</td><td>5.2</td><td>6.1234</td></tr></table></html>""".encode()
    session = Mock()
    session.get.return_value = Response(original)
    archive = ArchiveStore(tmp_path)
    extractor = DataExtractor(session, None, archive)
    bulletin = {"url": URL, "official_base_url": "https://tj.jiangsu.gov.cn",
                "region": "江苏省", "year": 2024, "type": "统计公报", "publish_date": "2025-03-01"}
    output = extractor.extract_data(bulletin)
    assert output["data"][INCOME] == 61234
    evidence = output["evidence"][INCOME]
    assert evidence["statistical_year"] == 2024
    assert evidence["raw_unit"] == "万元"
    assert evidence["locator"]["column"] == 3
    assert archive.resolve(evidence["archive_sha256"]).read_bytes() == original
    assert not Path(evidence["archive_path"]).is_absolute()
    assert extractor.extract_data(bulletin)["data"] == output["data"]
    assert session.get.call_count == 1


def test_missing_target_dimensions_never_uses_publish_date(tmp_path):
    session = Mock()
    output = DataExtractor(session, None, ArchiveStore(tmp_path)).extract_data(
        {"url": URL, "official_base_url": URL, "publish_date": "2025-01-01"})
    assert output["data"] == {}
    assert "统计年度" in output["warnings"][0]
    session.get.assert_not_called()


def test_source_disagreement_even_one_yuan_keeps_both_without_amount():
    record = DataRecord("江苏省", 2024)
    pipeline = collector()
    pipeline._merge_result(record, result(60000))
    pipeline._merge_result(record, result(60001))
    pipeline._merge_result(record, result(60000))
    assert record.urban_disposable_income is None
    assert len(record.candidates[INCOME]) == 3
    assert {item["value"] for item in record.evidence[INCOME]["conflicts"]} == {60000, 60001}
    json.dumps(record.to_dict())


def test_internal_conflict_cannot_be_replaced_by_another_source():
    conflict = result()
    conflict["data"] = {}
    conflict["evidence"][INCOME]["conflicts"] = [{"value": 60000}, {"value": 61000}]
    record = DataRecord("江苏省", 2024)
    pipeline = collector()
    pipeline._merge_result(record, conflict)
    assert len(record.evidence[INCOME]["conflicts"]) == 2
    pipeline._merge_result(record, result())
    assert record.urban_disposable_income is None


@pytest.mark.parametrize("change", [{"region": "浙江省"}, {"statistical_year": 2023}, {"unit": "万元"},
                                    {"validation_status": "pending"}])
def test_dimension_mismatch_cannot_enter_record(change):
    output = result()
    output["evidence"][INCOME].update(change)
    record = DataRecord("江苏省", 2024)
    collector()._merge_result(record, output)
    assert record.urban_disposable_income is None
    assert record.extraction_warnings


def test_nbs_only_completion_does_not_claim_wage_coverage():
    output = result()
    output["data"][CONSUMPTION] = 40000
    output["evidence"][CONSUMPTION] = {**output["evidence"][INCOME], "raw_value": "40000"}
    pipeline = collector()
    pipeline.nbs = Mock()
    pipeline.nbs.collect.return_value = output
    record = pipeline.collect_record("江苏省", 2024)
    assert record.status == "已提取"
    assert record.urban_non_private_wage is None
    pipeline.status.add_record(record)
    assert pipeline.status.is_completed("江苏省", 2024, extraction_version=f"{COLLECTION_VERSION}:nbs")
    assert not pipeline.status.is_completed("江苏省", 2024, extraction_version=f"{COLLECTION_VERSION}:auto")


def test_candidate_export_preserves_uncertainty_and_formula_escaping(tmp_path):
    pipeline = collector()
    output = result()
    candidate = output["candidates"][INCOME][0]
    candidate.update(raw_value="=1+1", validation_status="pending", review_reasons=["missing_unit"])
    pipeline.status.records = [DataRecord("江苏省", 2024, candidates={INCOME: [candidate]})]
    path = tmp_path / "result.xlsx"
    pipeline.export_to_excel(str(path))
    book = load_workbook(path, data_only=False)
    assert set(book.sheetnames) == {"统计数据", "字段溯源", "候选复核"}
    rows = list(book["候选复核"].values)
    candidate_row = dict(zip(rows[0], rows[1]))
    assert candidate_row["原始数值"] == "'=1+1"
    assert candidate_row["抽取校验"] == "pending"
    book.close()


def test_archive_is_immutable_and_rejects_traversal_and_corruption(tmp_path):
    archive = ArchiveStore(tmp_path)
    saved = archive.save(b"official", source_url=URL, content_type="text/html")
    digest = saved["archive_sha256"]
    assert archive.resolve("../" + digest) is None
    archive.resolve(digest).write_bytes(b"changed")
    assert archive.resolve(digest) is None
    with pytest.raises(ValueError, match="哈希"):
        archive.save(b"official", source_url=URL, content_type="text/html")


def test_post_redirect_is_not_followed():
    response = Response(b"")
    response.status_code = 307
    response.headers = {"Location": "https://evil.test/read"}
    session = Mock()
    session.post.return_value = response
    with pytest.raises(ValueError, match="POST"):
        fetch_official_content(session, URL, URL, method="POST", json_body={"year": 2024})
    assert session.post.call_count == 1
    session.get.assert_not_called()


def test_versioned_import_checks_dimensions_and_keeps_candidates(tmp_path):
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    sessions = create_session_factory(engine)
    record = DataRecord("江苏省", 2024, extraction_version="2026.2:nbs", status="已提取")
    collector()._merge_result(record, result())
    path = tmp_path / "status.json"
    path.write_text(json.dumps({"records": [record.to_dict()]}), encoding="utf-8")
    with sessions() as session:
        assert import_collector_status(session, path) == 1
        standard = session.query(StatisticalStandardModel).one()
        assert standard.status == "pending"
        assert len(standard.source_record["candidates"]) == 1
        record.evidence[INCOME]["statistical_year"] = 2023
        path.write_text(json.dumps({"records": [record.to_dict()]}), encoding="utf-8")
        with pytest.raises(ValueError, match="维度"):
            import_collector_status(session, path)


def test_new_conflict_invalidates_old_pending_value_and_clear_recollection_restores_pending(tmp_path):
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    sessions = create_session_factory(engine)
    pipeline = collector()
    record = DataRecord("江苏省", 2024, extraction_version="2026.2:auto", status="已提取")
    pipeline._merge_result(record, result(60000))
    path = tmp_path / "status.json"

    def write():
        path.write_text(json.dumps({"records": [record.to_dict()]}), encoding="utf-8")

    with sessions() as session:
        write()
        assert import_collector_status(session, path) == 1
        pipeline._merge_result(record, result(60001))
        write()
        assert import_collector_status(session, path) == 1
        standard = session.query(StatisticalStandardModel).one()
        assert standard.status == "invalidated"
        assert standard.source_record["evidence"]["conflicts"]
        assert standard.source_record["previous_observation"]["value"] == "60000.00"
        assert import_collector_status(session, path) == 1
        assert "previous_observation" not in standard.source_record["previous_observation"]["source_record"]
        record = DataRecord("江苏省", 2024, extraction_version="2026.2:auto", status="已提取")
        pipeline._merge_result(record, result(60001))
        write()
        assert import_collector_status(session, path) == 1
        assert standard.status == "pending"
        assert float(standard.value) == 60001
