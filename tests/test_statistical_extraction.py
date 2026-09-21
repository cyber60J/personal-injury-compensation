from __future__ import annotations

from io import BytesIO
import json
from unittest.mock import MagicMock, patch

import pytest
from openpyxl import Workbook

from personal_injury.collection.extraction import (
    CONSUMPTION, INCOME, NON_PRIVATE, ON_DUTY, PRIVATE, SERVICE, extract_document,
)


def html(body: str, *, year: int = 2024, region: str = "江苏省", title: str = "江苏省2024年统计公报"):
    return extract_document(body.encode("utf-8"), "text/html; charset=utf-8",
                            region=region, year=year, title=title,
                            source_url="https://tj.jiangsu.gov.cn/bulletin")


def test_prose_binds_year_and_keeps_old_candidate_without_selecting_it():
    result = html("<p>2023年城镇居民人均可支配收入为50000元，2024年城镇居民人均可支配收入为60000元，比上年增长20%。</p>")
    assert result["data"] == {INCOME: 60000}
    assert len(result["candidates"][INCOME]) == 2
    old = result["candidates"][INCOME][0]
    assert old["statistical_year"] == 2023
    assert "statistical_year_mismatch" in old["review_reasons"]
    assert result["evidence"][INCOME]["raw_value"] == "60000"


@pytest.mark.parametrize("columns,values", [("<th>2023年</th><th>2024年</th>", "<td>50000</td><td>60000</td>"),
                                            ("<th>2024年</th><th>2023年</th>", "<td>60000</td><td>50000</td>")])
def test_table_reads_year_column_instead_of_first_numeric_cell(columns, values):
    result = html(f"<table><caption>单位：元</caption><tr><th>指标</th>{columns}<th>增长率</th></tr><tr><td>{INCOME}</td>{values}<td>20</td></tr></table>")
    assert result["data"] == {INCOME: 60000}
    assert len(result["candidates"][INCOME]) == 2
    assert result["evidence"][INCOME]["locator"]["table"] == 1


def test_html_merged_headers_keep_population_period_and_units():
    result = html("""<table><tr><th rowspan="3">指标</th><th colspan="2">城镇居民</th><th rowspan="3">全体居民2024年</th></tr>
      <tr><th colspan="2">人均可支配收入（元）</th></tr><tr><th>2023年</th><th>2024年</th></tr>
      <tr><td>江苏省</td><td>50000</td><td>60000</td><td>40000</td></tr></table>""")
    assert result["data"] == {INCOME: 60000}
    evidence = result["evidence"][INCOME]
    assert evidence["locator"]["row"] == 4
    assert evidence["locator"]["column"] == 3


def test_wage_populations_are_never_interchangeable():
    result = html("""<p>城镇非私营单位就业人员年平均工资为120000元。</p>
      <p>城镇私营单位就业人员年平均工资为70000元。</p>
      <p>城镇非私营单位在岗职工年平均工资为130000元。</p>
      <p>城镇单位在岗职工年平均工资为125000元。</p>""")
    assert result["data"] == {NON_PRIVATE: 120000, PRIVATE: 70000, ON_DUTY: 125000}
    assert result["candidates"][ON_DUTY][0]["validation_status"] == "pending"
    assert result["candidates"][NON_PRIVATE][0]["scope"]["employment"] == "就业人员"


def test_missing_employment_scope_stays_pending():
    result = html("<p>城镇非私营单位年平均工资为120000元。</p>")
    assert result["data"] == {}
    assert "missing_employment_scope" in result["candidates"][NON_PRIVATE][0]["review_reasons"]


def test_industry_scope_comes_from_table_heading():
    result = html("""<h2>2024年城镇非私营单位就业人员分行业年平均工资</h2><p>单位：元</p>
      <table><tr><th>行业</th><th>2024年</th></tr><tr><td>居民服务、修理和其他服务业</td><td>65000</td></tr></table>""")
    assert result["data"] == {SERVICE: 65000}
    assert result["evidence"][SERVICE]["scope"]["ownership"] == "非私营"


def test_industry_without_unit_type_is_not_accepted():
    result = html("<p>居民服务、修理和其他服务业就业人员年平均工资为65000元。</p>")
    assert result["data"] == {}
    assert "service_requires_non_private_units" in result["candidates"][SERVICE][0]["review_reasons"]


def test_unit_conversion_preserves_original_value_and_formula():
    result = html("<p>城镇居民人均可支配收入为6.1234万元。</p>")
    assert result["data"] == {INCOME: 61234}
    evidence = result["evidence"][INCOME]
    assert evidence["raw_value"] == "6.1234"
    assert evidence["raw_unit"] == "万元"
    assert evidence["conversion"]["factor"] == 10000
    json.dumps(result, ensure_ascii=False)


def test_conflicting_totals_are_retained_and_neither_is_selected():
    result = html("<p>城镇居民人均消费支出为30000元。</p><p>城镇居民人均消费支出为31000元。</p>")
    assert result["data"] == {}
    assert len(result["evidence"][CONSUMPTION]["conflicts"]) == 2
    assert all(item["validation_status"] == "pending" for item in result["candidates"][CONSUMPTION])
    json.dumps(result)


def test_other_region_in_local_text_cannot_inherit_document_region():
    result = html("<p>浙江省城镇居民人均可支配收入为60000元。</p>")
    assert result["data"] == {}
    assert "region_mismatch" in result["candidates"][INCOME][0]["review_reasons"]


def test_other_region_title_overrides_search_title():
    result = html("<h1>浙江省2024年统计公报</h1><p>城镇居民人均可支配收入为60000元。</p>")
    assert result["data"] == {}
    assert "region_mismatch" in result["candidates"][INCOME][0]["review_reasons"]


def test_missing_unit_and_title_dimensions_are_reviewable():
    result = html("<table><tr><th>指标</th><th>2024年</th></tr><tr><td>城镇居民人均可支配收入</td><td>60000</td></tr></table>", title="统计公报")
    candidate = result["candidates"][INCOME][0]
    assert result["data"] == {}
    assert {"missing_or_ambiguous_unit", "missing_region"}.issubset(candidate["review_reasons"])
    assert candidate["conversion"] is None


def test_growth_rate_and_all_resident_values_are_not_candidates():
    result = html("<p>城镇居民人均可支配收入为5.2%。</p><p>全体居民人均可支配收入为40000元。</p><p>农村居民人均可支配收入为20000元。</p>")
    assert result["data"] == {}
    assert result["candidates"] == {}


def test_previous_year_word_cannot_inherit_target_year():
    result = html("<p>上年城镇居民人均可支配收入为50000元。</p>")
    assert result["data"] == {}
    assert result["candidates"][INCOME][0]["statistical_year"] == 2023


def test_xlsx_merged_headers_and_exact_source_cell():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "居民收入"
    sheet.append(["江苏省2024年统计公报"])
    sheet.merge_cells("A1:C1")
    sheet.append(["指标", "城镇居民", None])
    sheet.merge_cells("B2:C2")
    sheet.append([None, "人均可支配收入（万元）", None])
    sheet.merge_cells("B3:C3")
    sheet.append([None, "2023年", "2024年"])
    sheet.append(["江苏省", 5.0, 6.0])
    output = BytesIO()
    workbook.save(output)
    result = extract_document(output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", region="江苏省", year=2024)
    assert result["data"] == {INCOME: 60000}
    locator = result["evidence"][INCOME]["locator"]
    assert {key: locator[key] for key in ("type", "sheet", "cell")} == {"type": "xlsx_cell", "sheet": "居民收入", "cell": "C5"}
    assert "2024年" in locator["column_label"]


def test_empty_pdf_page_reports_ocr_instead_of_guessing():
    page = MagicMock()
    page.extract_text.return_value = ""
    document = MagicMock()
    document.__enter__.return_value.pages = [page]
    with patch("pdfplumber.open", return_value=document):
        result = extract_document(b"%PDF-fake", "application/pdf", region="江苏省", year=2024)
    assert result["data"] == {}
    assert result["warnings"] == ["needs_ocr: page 1"]
    page.find_tables.assert_not_called()


def test_pdf_text_retains_page_and_bounding_box():
    page = MagicMock()
    page.extract_text.return_value = "城镇居民人均可支配收入为60000元。"
    page.find_tables.return_value = []
    page.extract_words.return_value = [{"text": "城镇居民人均可支配收入为60000元。", "x0": 10, "x1": 300, "top": 20, "bottom": 35}]
    document = MagicMock()
    document.__enter__.return_value.pages = [page]
    with patch("pdfplumber.open", return_value=document):
        result = extract_document(b"%PDF-fake", "application/pdf", region="江苏省", year=2024, title="江苏省2024年统计公报")
    assert result["data"] == {INCOME: 60000}
    assert result["evidence"][INCOME]["locator"] == {"type": "pdf_text", "page": 1, "bbox": [10, 20, 300, 35]}


def test_input_limits_and_unsupported_formats_are_explicit():
    result = extract_document(b"binary", "application/octet-stream", region="江苏省", year=2024)
    assert result["warnings"] == ["unsupported_document_format"]
    result = html("<table><tr><td rowspan='1000000'>测试</td></tr></table>")
    assert "table_span_limit" in result["warnings"][0]


@pytest.mark.parametrize("body", ["全国，城镇居民人均可支配收入为54188元。", "南京市的城镇居民人均可支配收入为80673元。",
                                  "2024年南京市城镇居民人均可支配收入为80673元。", "上年，城镇居民人均可支配收入为60000元。",
                                  "睢宁县城镇居民人均可支配收入为40000元。", "南京市江宁区城镇居民人均可支配收入为70000元。"])
def test_sentence_dimensions_cannot_be_lost_at_a_comma(body):
    result = html(f"<p>{body}</p>")
    assert result["data"] == {}
    assert result["candidates"][INCOME][0]["validation_status"] == "pending"


@pytest.mark.parametrize("body", ["制造业城镇非私营单位就业人员年平均工资为120000元。", "其中，制造业就业人员年平均工资为120000元。",
                                  "制造业，城镇非私营单位就业人员年平均工资为120000元。"])
def test_specific_industries_cannot_be_used_as_all_industry_wages(body):
    result = html(f"<p>{body}</p>", title="江苏省2024年城镇非私营单位就业人员年平均工资")
    assert result["data"] == {}
    assert "non_total_industry_scope" in result["candidates"][NON_PRIVATE][0]["review_reasons"]


def test_pdf_own_heading_overrides_wrong_search_title():
    page = MagicMock()
    page.extract_text.return_value = "江苏省2023年统计公报\n城镇居民人均可支配收入为60000元。"
    page.find_tables.return_value = []
    page.extract_words.return_value = [{"text": "城镇居民人均可支配收入为60000元。", "x0": 10, "x1": 300, "top": 20, "bottom": 35}]
    document = MagicMock()
    document.__enter__.return_value.pages = [page]
    with patch("pdfplumber.open", return_value=document):
        result = extract_document(b"%PDF-fake", "application/pdf", region="江苏省", year=2024, title="江苏省2024年统计公报")
    assert result["data"] == {}
    assert result["candidates"][INCOME][0]["statistical_year"] == 2023


def test_corrupt_pdf_returns_explicit_warning():
    result = extract_document(b"%PDF-corrupt", "application/pdf", region="江苏省", year=2024)
    assert result["data"] == {}
    assert result["warnings"][0].startswith("document_extraction_failed:")


def test_footnotes_and_headers_are_preserved_but_not_used_as_amounts():
    result = html("<table><caption>单位：元</caption><tr><th>指标</th><th>2024年</th></tr><tr><td>城镇居民人均可支配收入</td><td>60000</td></tr></table><p>注：另表2023年数据为50000元。</p>")
    assert result["data"] == {INCOME: 60000}
    locator = result["evidence"][INCOME]["locator"]
    assert locator["row_label"] == INCOME
    assert locator["headers"] == [INCOME, "2024年"]
    assert locator["footnotes"] == ["注：另表2023年数据为50000元。"]
